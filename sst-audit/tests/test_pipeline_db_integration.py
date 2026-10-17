"""Интеграционные тесты сквозного пайплайна main.py с SQLite как SSOT.

Проверяется полный цикл: нормализованные DTO -> run_pipeline() -> записи в
sessions / rd_items / smeta_items / match_records / rule_evaluations,
сшивка отчета через v_audit_export и фильтрация очереди HITL.
"""

from collections.abc import Iterator
from decimal import Decimal
import io
import json
from pathlib import Path
import sqlite3

from openpyxl import Workbook, load_workbook
import pytest
from rich.console import Console

import db
import exporter
import main as main_module
from main import (
    AUTO_MATCH_THRESHOLD,
    MIN_MATCH_THRESHOLD,
    SessionConfig,
    match_items,
    run_pipeline,
)
from models import CellPointer, NormalizedItem, UnitClass

PROJECT_NAME = "ЖК Север, корпус 3"
SESSION_ID = "it-session-001"

TITLE_CABLE = "Кабель силовой ВВГнг(A)-LS 3х1.5"
CODE_CABLE = "ФССЦ-21.1.06.04-0012"
TITLE_VALVE = "Задвижка клиновая стальная фланцевая 30с41нж DN50 PN16"
CODE_VALVE = "30с41нж"
TITLE_PUMP = "Насос центробежный консольный"
CODE_PUMP = "PUMP-1"


def _pointer(doc_type: str, row_idx: int) -> CellPointer:
    """Создает аудиторский след ячейки для тестовых DTO."""
    if doc_type == "SPECIFICATION":
        return CellPointer(
            doc_type=doc_type,
            file_name="rd.pdf",
            sheet_name="Page 1",
            row_idx=row_idx,
            col_letter="D",
            coordinate=f"D{row_idx}",
        )
    return CellPointer(
        doc_type=doc_type,
        file_name="smeta.xlsx",
        sheet_name="Смета",
        row_idx=row_idx,
        col_letter="B",
        coordinate=f"B{row_idx}",
    )


def _rd_item(
    raw_id: str,
    title: str,
    code: str | None,
    unit: UnitClass,
    volume: Decimal,
    row_idx: int,
    raw_unit: str | None = None,
) -> NormalizedItem:
    """Позиция рабочей документации (спецификация)."""
    return NormalizedItem(
        raw_id=raw_id,
        code=code,
        raw_title=title,
        clean_title=title.lower(),
        raw_unit=raw_unit or unit.value,
        canonical_unit=unit,
        unit_multiplier=Decimal("1.0"),
        volume=volume,
        dimensions={},
        context_section="Default_Section",
        pointer=_pointer("SPECIFICATION", row_idx),
    )


def _smeta_item(
    raw_id: str,
    title: str,
    code: str | None,
    unit: UnitClass,
    volume: Decimal,
    row_idx: int,
    raw_unit: str | None = None,
    unit_price: Decimal = Decimal("0"),
    total_cost: Decimal = Decimal("0"),
) -> NormalizedItem:
    """Позиция сметы (материальный ресурс)."""
    return NormalizedItem(
        raw_id=raw_id,
        code=code,
        raw_title=title,
        clean_title=title.lower(),
        raw_unit=raw_unit or unit.value,
        canonical_unit=unit,
        unit_multiplier=Decimal("1.0"),
        volume=volume,
        dimensions={},
        context_section="Default_Section",
        pointer=_pointer("LSR", row_idx),
        is_resource=True,
        parent_id="LSR_1",
        unit_price=unit_price,
        total_cost=total_cost,
    )


def _cable_pair() -> tuple[NormalizedItem, NormalizedItem]:
    """Пара «кабель»: полное совпадение шифра и объема."""
    rd_item = _rd_item(
        "SPEC_1",
        TITLE_CABLE,
        CODE_CABLE,
        UnitClass.LENGTH,
        Decimal("150.0"),
        12,
        raw_unit="100 м",
    )
    smeta_item = _smeta_item(
        "LSR_10",
        TITLE_CABLE,
        CODE_CABLE,
        UnitClass.LENGTH,
        Decimal("150.0"),
        40,
        raw_unit="100 м",
        unit_price=Decimal("2500.00"),
        total_cost=Decimal("375000.00"),
    )
    return rd_item, smeta_item


def _make_two_section_smeta(path: Path) -> Path:
    """Создает двухлистовую смету для пакетного интеграционного теста."""
    workbook = Workbook()
    aps = workbook.active
    aps.title = "ЛС 02-01-08"
    aps["A8"] = "Автоматическая пожарная сигнализация (АПС)"

    sot = workbook.create_sheet("ЛС 02-01-14")
    sot["A9"] = "Система охранного телевидения (СОТ)"
    workbook.save(path)
    return path


def _make_batch_dir(tmp_path: Path) -> Path:
    """Создает папку с двумя PDF-томами, имена которых содержат разделы."""
    rd_dir = tmp_path / "rd"
    rd_dir.mkdir()
    (rd_dir / "Том 1 (АПС).pdf").write_bytes(b"")
    (rd_dir / "Том 2 (СОТ).pdf").write_bytes(b"")
    return rd_dir


def _three_item_dataset() -> tuple[list[NormalizedItem], list[NormalizedItem]]:
    """2 автоматические пары + 1 позиция РД без пары в смете."""
    rd_cable, smeta_cable = _cable_pair()
    rd_valve = _rd_item(
        "SPEC_2",
        TITLE_VALVE,
        CODE_VALVE,
        UnitClass.COUNT,
        Decimal("8.0"),
        15,
    )
    smeta_valve = _smeta_item(
        "LSR_20",
        TITLE_VALVE,
        CODE_VALVE,
        UnitClass.COUNT,
        Decimal("8.0"),
        45,
        unit_price=Decimal("12000.00"),
        total_cost=Decimal("96000.00"),
    )
    rd_pump = _rd_item(
        "SPEC_3",
        TITLE_PUMP,
        CODE_PUMP,
        UnitClass.COUNT,
        Decimal("2.0"),
        18,
    )
    return [rd_cable, rd_valve, rd_pump], [smeta_cable, smeta_valve]


def _config(session_id: str = SESSION_ID) -> SessionConfig:
    """Конфигурация тестовой сессии (пути не читаются: DTO переданы явно)."""
    return SessionConfig(
        project_name=PROJECT_NAME,
        rd_path="rd.pdf",
        smeta_path="smeta.xlsx",
        session_id=session_id,
    )


@pytest.fixture()
def conn(tmp_path: Path) -> Iterator[sqlite3.Connection]:
    """Реальная файловая БД SQLite (WAL), изолированная на тест."""
    connection = db.get_connection(tmp_path / "sst_audit_it.db")
    try:
        yield connection
    finally:
        connection.close()


def _count(connection: sqlite3.Connection, table: str) -> int:
    return int(
        connection.execute(f"SELECT COUNT(*) AS c FROM {table}").fetchone()["c"]
    )


def test_full_pipeline_persists_all_entities(conn: sqlite3.Connection) -> None:
    """Сквозной прогон создает записи во всех таблицах SSOT."""
    rd_items, smeta_items = _three_item_dataset()

    result = run_pipeline(conn, _config(), rd_items=rd_items, smeta_items=smeta_items)

    assert result.session_id == SESSION_ID
    assert result.rd_saved == 3
    assert result.smeta_saved == 2
    # 2 решения по смете + 1 обратный проход по неосмеченной позиции РД
    assert result.matches_saved == 3
    assert result.evaluations_saved == 3
    assert result.auto_matches == 2
    assert result.hitl_pending == 1
    assert result.session_status == "MATCHED"

    assert _count(conn, "sessions") == 1
    assert _count(conn, "rd_items") == 3
    assert _count(conn, "smeta_items") == 2
    assert _count(conn, "match_records") == 3
    assert _count(conn, "rule_evaluations") == 3

    session = conn.execute(
        "SELECT * FROM sessions WHERE id = ?", (SESSION_ID,)
    ).fetchone()
    assert session is not None
    assert session["project_name"] == PROJECT_NAME
    assert session["status"] == "MATCHED"

    statuses = {
        row["status"]
        for row in conn.execute(
            "SELECT status FROM match_records WHERE session_id = ?", (SESSION_ID,)
        )
    }
    assert statuses == {"AUTO", "HITL_PENDING"}

    rules = {
        row["rule_name"]
        for row in conn.execute("SELECT rule_name FROM rule_evaluations")
    }
    assert "RD_NOT_IN_SMETA" in rules
    assert "VOLUME_WITHIN_TOLERANCE" in rules


def test_main_exports_excel_without_interactive_review(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """CLI ``--no-review`` автоматически формирует итоговую ведомость XLSX."""
    rd_items, smeta_items = _three_item_dataset()

    class _StubPDFParser:
        def __init__(self, _path: str) -> None:
            pass

        def extract_items(self) -> list[NormalizedItem]:
            return rd_items

    class _StubLSRParser:
        def __init__(self, _path: str) -> None:
            pass

        def parse(self) -> list[NormalizedItem]:
            return smeta_items

    monkeypatch.setattr(main_module, "PDFSpecParser", _StubPDFParser)
    monkeypatch.setattr(main_module, "LSRParser", _StubLSRParser)

    rd_path = tmp_path / "rd.pdf"
    smeta_path = tmp_path / "smeta.xlsx"
    rd_path.write_bytes(b"")
    smeta_path.write_bytes(b"")
    target = tmp_path / "audit.xlsx"

    code = main_module.main(
        [
            "--project",
            PROJECT_NAME,
            "--rd",
            str(rd_path),
            "--smeta",
            str(smeta_path),
            "--db",
            str(tmp_path / "cli.db"),
            "--session-id",
            "it-cli-export",
            "--out",
            str(target),
            "--no-review",
        ]
    )

    assert code == 0
    assert target.is_file()
    workbook = load_workbook(target, data_only=False)
    assert exporter.SHEET_TITLE in workbook.sheetnames


def test_pipeline_filters_noise_below_min_threshold(conn: sqlite3.Connection) -> None:
    """Шум (<40%) сохраняется как NO_MATCH и не создает карточку HITL."""
    rd_item = _rd_item(
        "SPEC_NOISE_PIPE",
        TITLE_CABLE,
        CODE_CABLE,
        UnitClass.LENGTH,
        Decimal("10.0"),
        70,
    )
    noise_smeta = _smeta_item(
        "LSR_NOISE_PIPE",
        "Кирпич керамический полнотелый одинарный",
        None,
        UnitClass.LENGTH,
        Decimal("100.0"),
        71,
    )

    result = run_pipeline(
        conn,
        _config("it-session-noise"),
        rd_items=[rd_item],
        smeta_items=[noise_smeta],
    )

    assert result.no_match == 1
    assert result.hitl_pending == 0
    assert result.auto_matches == 0

    match = conn.execute(
        "SELECT * FROM match_records WHERE session_id = ?",
        ("it-session-noise",),
    ).fetchone()
    assert match is not None
    assert match["status"] == "AUTO"
    assert match["match_method"] == "NO_MATCH"
    assert match["rd_item_id"] is None
    assert match["confidence_score"] == 0
    assert db.get_pending_reviews(conn, "it-session-noise") == []


def test_pipeline_handles_empty_rd_gracefully(conn: sqlite3.Connection) -> None:
    """Нераспознанный РД не роняет пайплайн: все сметы фиксируются в SMETA_NOT_IN_RD."""
    _, smeta_item = _cable_pair()

    result = run_pipeline(
        conn,
        _config("it-session-empty-rd"),
        rd_items=[],
        smeta_items=[smeta_item],
    )

    assert result.rd_saved == 0
    assert result.smeta_saved == 1
    assert result.matches_saved == 1
    assert result.auto_matches == 0
    assert result.hitl_pending == 0
    assert result.no_match == 1
    assert result.session_status == "REVIEWED"

    match = conn.execute(
        "SELECT * FROM match_records WHERE session_id = ?",
        ("it-session-empty-rd",),
    ).fetchone()
    assert match is not None
    assert match["status"] == "AUTO"
    assert match["rd_item_id"] is None
    assert match["confidence_score"] == 0

    rule = conn.execute(
        "SELECT rule_name FROM rule_evaluations WHERE match_record_id = ?",
        (match["id"],),
    ).fetchone()
    assert rule is not None
    assert rule["rule_name"] == "SMETA_NOT_IN_RD"


def test_audit_export_view_stitches_rows(conn: sqlite3.Connection) -> None:
    """v_audit_export возвращает корректно сшитые строки RD + смета + связь."""
    rd_items, smeta_items = _three_item_dataset()
    run_pipeline(conn, _config(), rd_items=rd_items, smeta_items=smeta_items)

    rows = conn.execute(
        "SELECT * FROM v_audit_export WHERE session_id = ?", (SESSION_ID,)
    ).fetchall()
    assert len(rows) == 3

    by_rd_code = {row["rd_code"]: row for row in rows}

    cable = by_rd_code[CODE_CABLE]
    assert cable["project_name"] == PROJECT_NAME
    assert cable["smeta_code"] == CODE_CABLE
    assert cable["rd_quantity"] == pytest.approx(150.0)
    assert cable["smeta_quantity"] == pytest.approx(150.0)
    assert cable["unit_price"] == pytest.approx(2500.0)
    assert cable["total_cost"] == pytest.approx(375000.0)
    assert cable["match_status"] == "AUTO"
    assert cable["rd_unit"] == "100 м"

    payload = json.loads(cable["rd_raw_payload"])
    assert payload["raw_id"] == "SPEC_1"
    assert payload["volume"] == "150.0"
    assert payload["canonical_unit"] == "м"
    assert payload["pointer"]["coordinate"] == "D12"

    pump = by_rd_code[CODE_PUMP]
    assert pump["smeta_code"] is None
    assert pump["match_status"] == "HITL_PENDING"
    assert pump["match_method"] == "REVERSE_PASS"


def test_pending_reviews_only_hitl_pending(conn: sqlite3.Connection) -> None:
    """get_pending_reviews отдает только связи со статусом HITL_PENDING."""
    rd_items, smeta_items = _three_item_dataset()
    run_pipeline(conn, _config(), rd_items=rd_items, smeta_items=smeta_items)

    pending = db.get_pending_reviews(conn, SESSION_ID)

    assert len(pending) == 1
    assert {row["status"] for row in pending} == {"HITL_PENDING"}
    assert pending[0]["rd_code"] == CODE_PUMP
    assert pending[0]["smeta_item_id"] is None

    db.record_expert_review(
        conn,
        int(pending[0]["match_id"]),
        "Инженер ПТО",
        "REJECTED",
        reason_code="NOT_IN_SCOPE",
        comment="Позиция не требует сверки со сметой",
    )
    assert db.get_pending_reviews(conn, SESSION_ID) == []


def test_session_reviewed_when_no_pending(conn: sqlite3.Connection) -> None:
    """Без очереди HITL сессия получает статус REVIEWED."""
    rd_cable, smeta_cable = _cable_pair()

    result = run_pipeline(
        conn,
        _config("it-session-reviewed"),
        rd_items=[rd_cable],
        smeta_items=[smeta_cable],
    )

    assert result.session_status == "REVIEWED"
    assert result.hitl_pending == 0
    assert result.matches_saved == 1
    assert result.evaluations_saved == 1
    assert db.get_pending_reviews(conn, "it-session-reviewed") == []


def test_ingest_rollback_keeps_database_clean(
    conn: sqlite3.Connection,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Сбой на этапе ingest откатывает сессию и все ранее вставленные позиции."""
    rd_cable, smeta_cable = _cable_pair()

    def _fail_on_smeta(*args: object, **kwargs: object) -> int:
        raise RuntimeError("disk full")

    monkeypatch.setattr(db, "save_smeta_items", _fail_on_smeta)

    with pytest.raises(RuntimeError, match="disk full"):
        run_pipeline(
            conn,
            _config("it-session-rollback"),
            rd_items=[rd_cable],
            smeta_items=[smeta_cable],
        )

    for table in (
        "sessions",
        "rd_items",
        "smeta_items",
        "match_records",
        "rule_evaluations",
    ):
        assert _count(conn, table) == 0, f"Таблица {table} не очищена rollback-ом"


def test_match_items_marks_low_confidence_and_ambiguity() -> None:
    """Матчер классифицирует низкую уверенность и неоднозначность как HITL."""
    weak_smeta = _smeta_item(
        "LSR_W",
        "Кабель силовой",
        None,
        UnitClass.LENGTH,
        Decimal("10.0"),
        50,
    )
    control_rd = _rd_item(
        "SPEC_W",
        "Кабель контрольный КВВГ 3х1.5",
        None,
        UnitClass.LENGTH,
        Decimal("10.0"),
        20,
    )

    weak_decisions = match_items([control_rd], [weak_smeta])
    assert len(weak_decisions) == 1
    assert weak_decisions[0].rd_item is control_rd
    assert weak_decisions[0].confidence < AUTO_MATCH_THRESHOLD
    assert weak_decisions[0].status == "HITL_PENDING"

    ambiguous_smeta = _smeta_item(
        "LSR_A",
        "Кабель силовой ВВГнг 3х1.5",
        None,
        UnitClass.LENGTH,
        Decimal("10.0"),
        51,
    )
    exact_rd = _rd_item(
        "SPEC_A",
        "Кабель силовой ВВГнг 3х1.5",
        None,
        UnitClass.LENGTH,
        Decimal("10.0"),
        21,
    )
    similar_rd = _rd_item(
        "SPEC_B",
        "Кабель силовой ВВГнг 3х2.5",
        None,
        UnitClass.LENGTH,
        Decimal("10.0"),
        22,
    )

    ambiguous_decisions = match_items([exact_rd, similar_rd], [ambiguous_smeta])
    top = ambiguous_decisions[0]
    assert top.rd_item is exact_rd
    assert top.confidence >= AUTO_MATCH_THRESHOLD
    assert len(top.candidates) == 2
    assert top.status == "HITL_PENDING"


def test_match_items_below_min_threshold_is_no_match() -> None:
    """Сходство ниже 40% не попадает в HITL и не создает шумовую карточку."""
    rd_item = _rd_item(
        "SPEC_NOISE",
        TITLE_CABLE,
        CODE_CABLE,
        UnitClass.LENGTH,
        Decimal("10.0"),
        40,
    )
    noise_smeta = _smeta_item(
        "LSR_NOISE",
        "Кирпич керамический полнотелый одинарный",
        None,
        UnitClass.LENGTH,
        Decimal("100.0"),
        41,
    )

    decisions = match_items([rd_item], [noise_smeta])

    assert len(decisions) == 1
    decision = decisions[0]
    assert decision.status == "AUTO"
    assert decision.method == "NO_MATCH"
    assert decision.rd_item is None
    assert decision.confidence == Decimal("0")
    assert decision.candidates == ()
    assert decision.confidence < MIN_MATCH_THRESHOLD


def test_interactive_hitl_updates_link_and_rules(
    conn: sqlite3.Connection,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Выбор второго кандидата в HITL обновляет связь и оценку правил."""
    smeta = _smeta_item(
        "LSR_INT",
        "Кабель силовой ВВГнг 3х1.5",
        None,
        UnitClass.LENGTH,
        Decimal("10.0"),
        60,
        unit_price=Decimal("100.00"),
        total_cost=Decimal("1000.00"),
    )
    exact = _rd_item(
        "SPEC_X",
        "Кабель силовой ВВГнг 3х1.5",
        "ART-1.5",
        UnitClass.LENGTH,
        Decimal("10.0"),
        30,
    )
    similar = _rd_item(
        "SPEC_Y",
        "Кабель силовой ВВГнг 3х2.5",
        "ART-2.5",
        UnitClass.LENGTH,
        Decimal("20.0"),
        31,
    )

    # Первая карточка: выбор второго кандидата; вторая (обратный проход) —
    # пропуск, чтобы очередь сохранила MATCHED.
    answers = iter(["2", "S"])
    monkeypatch.setattr("builtins.input", lambda _prompt: next(answers))

    result = run_pipeline(
        conn,
        _config("it-session-hitl"),
        rd_items=[exact, similar],
        smeta_items=[smeta],
        interactive=True,
        reviewer="Инженер ПТО",
    )

    # Обратный проход по второму кандидату остается в очереди HITL.
    assert result.session_status == "MATCHED"

    match = conn.execute(
        "SELECT * FROM match_records WHERE session_id = ? ORDER BY id LIMIT 1",
        ("it-session-hitl",),
    ).fetchone()
    assert match is not None
    assert match["status"] == "RESOLVED"

    selected_id = conn.execute(
        "SELECT id FROM rd_items WHERE code = 'ART-2.5'"
    ).fetchone()["id"]
    assert match["rd_item_id"] == selected_id
    assert match["confidence_score"] == pytest.approx(96.153846, rel=1e-4)

    review = conn.execute(
        "SELECT * FROM expert_reviews WHERE match_record_id = ?",
        (match["id"],),
    ).fetchone()
    assert review is not None
    assert review["decision"] == "APPROVED"
    assert review["reviewer_name"] == "Инженер ПТО"

    rule = conn.execute(
        "SELECT * FROM rule_evaluations WHERE match_record_id = ?",
        (match["id"],),
    ).fetchone()
    assert rule is not None
    assert rule["rule_name"] == "VOLUME_DELTA_SUSPECT"
    assert rule["is_violation"] == 1
    # Дельта рассчитана по выбранному кандидату (20 x 1.06), а не по первому.
    assert rule["delta_volume"] == pytest.approx(-11.2)


def test_batch_pipeline_headless(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Ночной прогон создает по сессии на раздел и не блокирует ввод."""
    rd_dir = _make_batch_dir(tmp_path)
    smeta_path = _make_two_section_smeta(tmp_path / "smeta.xlsx")
    db_path = tmp_path / "batch.db"
    target = tmp_path / "batch_report.xlsx"

    class _StubPDFParser:
        def __init__(self, _path: str) -> None:
            self._rd_item = _cable_pair()[0]

        def extract_items(self) -> list[NormalizedItem]:
            return [self._rd_item]

    class _StubLSRParser:
        def __init__(self, _path: str) -> None:
            pass

        def parse(self, sheet_name: str | None = None) -> list[NormalizedItem]:
            assert sheet_name in {"ЛС 02-01-08", "ЛС 02-01-14"}
            return [_cable_pair()[1]]

    def _forbidden_input(_prompt: str = "") -> str:
        raise AssertionError("Headless-режим не должен запрашивать ввод оператора")

    monkeypatch.setattr(main_module, "PDFSpecParser", _StubPDFParser)
    monkeypatch.setattr(main_module, "LSRParser", _StubLSRParser)
    monkeypatch.setattr("builtins.input", _forbidden_input)

    code = main_module.main(
        [
            "--project",
            PROJECT_NAME,
            "--batch-dir",
            str(rd_dir),
            "--smeta",
            str(smeta_path),
            "--db",
            str(db_path),
            "--out",
            str(target),
            "--headless",
        ]
    )

    assert code == 0
    conn = db.get_connection(db_path)
    try:
        sessions = conn.execute(
            "SELECT project_name, smeta_sheet, rd_filename FROM sessions "
            "ORDER BY smeta_sheet"
        ).fetchall()
    finally:
        conn.close()

    assert [row["smeta_sheet"] for row in sessions] == [
        "ЛС 02-01-08",
        "ЛС 02-01-14",
    ]
    assert {row["project_name"] for row in sessions} == {PROJECT_NAME}
    assert len({row["rd_filename"] for row in sessions}) == 2

    workbook = load_workbook(target, data_only=False)
    assert workbook.sheetnames == [
        exporter.DASHBOARD_TITLE,
        "АПС (ЛС 02-01-08)",
        "СОТ (ЛС 02-01-14)",
    ]


def test_batch_pipeline_continues_after_section_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Исключение в одном томе логируется, но не прерывает остальные разделы."""
    rd_dir = _make_batch_dir(tmp_path)
    smeta_path = _make_two_section_smeta(tmp_path / "smeta.xlsx")
    db_path = tmp_path / "batch-errors.db"
    error_log = tmp_path / "batch_errors.log"
    target = tmp_path / "batch_report.xlsx"

    class _FlakyPDFParser:
        def __init__(self, path: str) -> None:
            if "СОТ" in Path(path).name:
                raise RuntimeError("поврежденный PDF")
            self._rd_item = _cable_pair()[0]

        def extract_items(self) -> list[NormalizedItem]:
            return [self._rd_item]

    class _StubLSRParser:
        def __init__(self, _path: str) -> None:
            pass

        def parse(self, sheet_name: str | None = None) -> list[NormalizedItem]:
            return [_cable_pair()[1]]

    monkeypatch.setattr(main_module, "PDFSpecParser", _FlakyPDFParser)
    monkeypatch.setattr(main_module, "LSRParser", _StubLSRParser)

    conn = db.get_connection(db_path)
    try:
        summary = main_module.run_batch_pipeline(
            conn,
            project_name=PROJECT_NAME,
            batch_dir=rd_dir,
            smeta_path=smeta_path,
            db_path=str(db_path),
            headless=True,
            output_path=target,
            error_log_path=error_log,
            console=Console(file=io.StringIO(), force_terminal=False),
        )
        sessions = conn.execute(
            "SELECT smeta_sheet FROM sessions ORDER BY smeta_sheet"
        ).fetchall()
    finally:
        conn.close()

    assert summary.total == 2
    assert summary.succeeded == 1
    assert summary.failed == 1
    assert [row["smeta_sheet"] for row in sessions] == ["ЛС 02-01-08"]
    assert target.is_file()
    assert "поврежденный PDF" in error_log.read_text(encoding="utf-8")
    assert "RuntimeError" in error_log.read_text(encoding="utf-8")
