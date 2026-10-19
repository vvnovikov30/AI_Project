"""Тесты автоматического сопоставления томов РД и листов многолистовой сметы."""

from pathlib import Path

from openpyxl import Workbook
import pytest

from dispatcher import (
    MAPPING_RULES,
    ProjectDispatcher,
    ROLE_SMETA,
    ROLE_VOR,
    classify_file_role,
    classify_smeta_tag,
    extract_section_code,
    pair_dual_directories,
    score_pair,
)


def _make_smeta(path: Path) -> Path:
    """Создает смету с двумя разделами ЛС и служебными листами."""
    workbook = Workbook()

    service = workbook.active
    service.title = "ССРСС"
    service["A8"] = "Сводный сметный расчет (АПС)"

    aps = workbook.create_sheet("ЛС 02-01-08")
    aps["A8"] = "Автоматическая пожарная сигнализация (АПС)"

    sot = workbook.create_sheet("ЛС 02-01-14")
    # Проверяем чтение заголовка из A9, если A8 пуст.
    sot["A9"] = "Система охранного телевидения (СОТ)"

    other = workbook.create_sheet("ОСР")
    other["A8"] = "Объектная смета (СОТ)"

    workbook.save(path)
    return path


def test_mapping_rules_cover_required_sections() -> None:
    """Словарь содержит все обязательные аббревиатуры разделов."""
    assert {
        "АПС",
        "СОУЭ",
        "РТС",
        "СКТ",
        "СКС",
        "СОТ",
        "СКУД",
        "СДС",
        "ЭОМ",
        "ОВ",
        "ВК",
        "АР",
    } <= set(MAPPING_RULES)


def test_match_volumes_maps_pdf_to_ls_sheet(tmp_path: Path) -> None:
    """PDF с аббревиатурой раздела привязывается к профильному листу ЛС."""
    rd_dir = tmp_path / "rd"
    rd_dir.mkdir()
    (rd_dir / "Том 1 (АПС) 08.04.2026.pdf").write_bytes(b"")
    (rd_dir / "Том 2 (СОТ) 08.04.2026.pdf").write_bytes(b"")
    (rd_dir / "Пояснительная записка.pdf").write_bytes(b"")
    (rd_dir / "readme.txt").write_text("not a pdf", encoding="utf-8")

    smeta_path = _make_smeta(tmp_path / "smeta.xlsx")
    pairs = ProjectDispatcher().match_volumes(rd_dir, smeta_path)

    assert [(pdf.name, sheet) for pdf, sheet, _score in pairs] == [
        ("Том 1 (АПС) 08.04.2026.pdf", "ЛС 02-01-08"),
        ("Том 2 (СОТ) 08.04.2026.pdf", "ЛС 02-01-14"),
    ]
    assert all(score >= ProjectDispatcher.MIN_CONFIDENCE for _, _, score in pairs)


def test_match_volumes_ignores_service_sheets_and_unrelated_pdf(
    tmp_path: Path,
) -> None:
    """Служебные листы и PDF без подтвержденного раздела не попадают в пары."""
    rd_dir = tmp_path / "rd"
    rd_dir.mkdir()
    (rd_dir / "Разное.pdf").write_bytes(b"")

    smeta_path = _make_smeta(tmp_path / "smeta.xlsx")
    sections = ProjectDispatcher().read_sections(smeta_path)

    assert [section.name for section in sections] == [
        "ЛС 02-01-08",
        "ЛС 02-01-14",
    ]
    assert ProjectDispatcher().match_volumes(rd_dir, smeta_path) == []


def test_score_pair_requires_section_evidence() -> None:
    """Высокая лексическая близость без подтверждения раздела не проходит порог."""
    no_evidence = score_pair("Том 5", "ЛС 02-01-14", "Система охраны")
    confirmed = score_pair(
        "Том 5 (СОТ)",
        "ЛС 02-01-14",
        "Система охранного телевидения (СОТ)",
    )

    assert no_evidence < ProjectDispatcher.MIN_CONFIDENCE
    assert confirmed >= ProjectDispatcher.MIN_CONFIDENCE
    assert confirmed > no_evidence


def test_duplicate_section_pdf_does_not_reuse_sheet(tmp_path: Path) -> None:
    """Один лист ЛС не может быть назначен двум разным PDF-томам."""
    rd_dir = tmp_path / "rd"
    rd_dir.mkdir()
    (rd_dir / "Том 1 (АПС).pdf").write_bytes(b"")
    (rd_dir / "Том 2 (АПС).pdf").write_bytes(b"")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "ЛС 02-01-08"
    worksheet["A8"] = "Автоматическая пожарная сигнализация (АПС)"
    smeta_path = tmp_path / "smeta.xlsx"
    workbook.save(smeta_path)

    pairs = ProjectDispatcher().match_volumes(rd_dir, smeta_path)

    assert len(pairs) == 1
    assert pairs[0][1] == "ЛС 02-01-08"


def test_match_volumes_rejects_missing_input(tmp_path: Path) -> None:
    """Отсутствие папки РД или сметы дает явную ошибку."""
    with pytest.raises(FileNotFoundError):
        ProjectDispatcher().match_volumes(tmp_path / "missing", tmp_path / "x.xlsx")

    rd_dir = tmp_path / "rd"
    rd_dir.mkdir()
    with pytest.raises(FileNotFoundError):
        ProjectDispatcher().match_volumes(rd_dir, tmp_path / "missing.xlsx")


def test_tag_classification_vkl_iskl() -> None:
    """Маркеры включения и исключения распознаются в именах ЛСР."""
    assert classify_smeta_tag("02-01-02_ЛСР_включение.xlsx") == "ВКЛ"
    assert classify_smeta_tag("02-01-02_доп.xlsx") == "ВКЛ"
    assert classify_smeta_tag("02-01-02_исключение.xlsx") == "ИСКЛ"
    assert classify_smeta_tag("02-01-02_снятие.pdf") == "ИСКЛ"
    assert classify_smeta_tag("02-01-02_ЛСР.xlsx") == "ОСН"
    assert extract_section_code("ВОР 02-05.pdf") == "02-05"


def test_dual_directory_pairing(tmp_path: Path) -> None:
    """Связки строятся по шифру раздела, СВОР приоритетнее ВОР."""
    rd_dir = tmp_path / "wor"
    smeta_dir = tmp_path / "lsr"
    rd_dir.mkdir()
    smeta_dir.mkdir()

    for name in (
        "02-01-01_ВОР.pdf",
        "02-01-03_ВОР.pdf",
        "02-01-03_СВОР.pdf",
        "Пояснительная записка.pdf",
    ):
        (rd_dir / name).write_bytes(b"")
    for name in (
        "02-01-01_ЛСР.xlsx",
        "02-01-01_ЛСР_включение.xlsx",
        "02-01-01_ЛСР_исключение.xlsx",
        "02-01-03_ЛСР_исключение.pdf",
        "~$02-01-01_ЛСР.xlsx",
    ):
        (smeta_dir / name).write_bytes(b"")

    pairs = ProjectDispatcher().match_dual_directories(rd_dir, smeta_dir)

    assert [
        (pair.section_code, pair.tag, pair.smeta_path.name) for pair in pairs
    ] == [
        ("02-01-01", "ОСН", "02-01-01_ЛСР.xlsx"),
        ("02-01-01", "ВКЛ", "02-01-01_ЛСР_включение.xlsx"),
        ("02-01-01", "ИСКЛ", "02-01-01_ЛСР_исключение.xlsx"),
        ("02-01-03", "ИСКЛ", "02-01-03_ЛСР_исключение.pdf"),
    ]
    assert all(pair.rd_path.name == "02-01-01_ВОР.pdf" for pair in pairs[:3])
    assert pairs[3].rd_path.name == "02-01-03_СВОР.pdf"
    assert pairs[0].sheet_name == "02-01-01 ОСН"
    assert pairs[1].session_id_prefix == "02-01-01_ВКЛ"
    assert pairs[0].session_id_prefix != pairs[1].session_id_prefix


def test_pair_dual_directories_is_pure_and_deduplicates() -> None:
    """Чистая функция не трогает диск и схлопывает дубли по (шифр, тег)."""
    rd_files = [Path("02-01-05_ВОР.pdf"), Path("02-01-05_СВОР.pdf")]
    smeta_files = [
        Path("02-01-05_ЛСР_исключение (1).xlsx"),
        Path("02-01-05_ЛСР_исключение.xlsx"),
        Path("02-01-05_ЛСР_исключение_копия.xlsx"),
        Path("02-01-06_ЛСР.xlsx"),
    ]

    pairs = pair_dual_directories(rd_files, smeta_files)

    assert len(pairs) == 1
    assert pairs[0].rd_path.name == "02-01-05_СВОР.pdf"
    assert pairs[0].smeta_path.name == "02-01-05_ЛСР_исключение.xlsx"


def test_single_pair_fallback_without_section_codes(tmp_path: Path) -> None:
    """Единственная ведомость и единственная смета связываются как 01-01/ОСН."""
    rd_dir = tmp_path / "rd"
    smeta_dir = tmp_path / "smeta"
    rd_dir.mkdir()
    smeta_dir.mkdir()

    rd_name = "Магазин_Анастасьевка_Дефектная_ведомость.xlsx"
    smeta_name = "Магазин_Анастасьевка_ЛСР_по_Методике_2020_РИМ.xlsx"
    (rd_dir / rd_name).write_bytes(b"")
    (smeta_dir / smeta_name).write_bytes(b"")

    pairs = ProjectDispatcher().match_dual_directories(rd_dir, smeta_dir)

    assert len(pairs) == 1
    pair = pairs[0]
    assert (pair.section_code, pair.tag) == ("01-01", "ОСН")
    assert pair.rd_path.name == rd_name
    assert pair.smeta_path.name == smeta_name
    assert pair.sheet_name == "01-01 ОСН"

    # Дефектная ведомость и смета в одной папке: дубли кандидатов схлопываются.
    flat_dir = tmp_path / "flat"
    flat_dir.mkdir()
    (flat_dir / rd_name).write_bytes(b"")
    (flat_dir / smeta_name).write_bytes(b"")

    same_dir_pairs = ProjectDispatcher().match_dual_directories(flat_dir, flat_dir)

    assert len(same_dir_pairs) == 1
    assert same_dir_pairs[0].section_code == "01-01"
    assert same_dir_pairs[0].rd_path.name == rd_name
    assert same_dir_pairs[0].smeta_path.name == smeta_name


def test_excel_vor_role_detection() -> None:
    """XLSX с маркерами ВОР отделяется от сметных расчетов в той же папке."""
    assert classify_file_role("Дефектная_ведомость_Магазин.xlsx") == ROLE_VOR
    assert classify_file_role("ВОР_ЭОМ.xlsx") == ROLE_VOR
    assert classify_file_role("СВОР_объект.xlsx") == ROLE_VOR

    assert classify_file_role("ЛСР_по_Методике_2020_РИМ.xlsx") == ROLE_SMETA
    assert classify_file_role("Смета_на_монтаж.xlsx") == ROLE_SMETA
    assert classify_file_role("ЛСР_ТЕР.xlsx") == ROLE_SMETA

    # Сметный PDF и PDF тома РД различаются по маркерам, lock-копии отсекаются.
    assert classify_file_role("02-01-03_ЛСР_исключение.pdf") == ROLE_SMETA
    assert classify_file_role("02-01-03_ВОР.pdf") == ROLE_VOR
    assert classify_file_role("~$ЛСР_РИМ.xlsx") is None
    assert classify_file_role("readme.txt") is None


def test_prefix_match_pairs_object_without_section_codes() -> None:
    """Без шифров пары связываются по общему корню имени объекта."""
    rd_files = [
        Path("Магазин_Анастасьевка_Дефектная_ведомость.xlsx"),
        Path("Школа_Дефектная_ведомость.xlsx"),
    ]
    smeta_files = [
        Path("Магазин_Анастасьевка_ЛСР_РИМ.xlsx"),
        Path("Школа_ЛСР.xlsx"),
    ]

    pairs = pair_dual_directories(rd_files, smeta_files)

    assert [
        (pair.section_code, pair.rd_path.name, pair.smeta_path.name)
        for pair in pairs
    ] == [
        (
            "01-01",
            "Магазин_Анастасьевка_Дефектная_ведомость.xlsx",
            "Магазин_Анастасьевка_ЛСР_РИМ.xlsx",
        ),
        (
            "01-02",
            "Школа_Дефектная_ведомость.xlsx",
            "Школа_ЛСР.xlsx",
        ),
    ]
