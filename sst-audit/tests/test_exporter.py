"""Тесты генератора ведомости аудита в Excel (exporter.py).

Проверяется структура книги, нативные формулы проверки, скрытая служебная
колонка и приоритет экспертного коэффициента ``adjusted_ratio``.
"""

from collections.abc import Iterator
from decimal import Decimal
from pathlib import Path
import sqlite3

from openpyxl import load_workbook
import pytest

import db
import exporter
from exporter import (
    DASHBOARD_TITLE,
    export_audit_to_excel,
    export_sessions_to_excel,
)

SESSION_ID = "export-session-001"
PROJECT_NAME = "ЖК Север, корпус 3"


@pytest.fixture()
def conn(tmp_path: Path) -> Iterator[sqlite3.Connection]:
    """Изолированная файловая БД SQLite с инициализированной схемой SSOT."""
    connection = db.get_connection(tmp_path / "exporter.db")
    db.init_db(connection)
    try:
        yield connection
    finally:
        connection.close()


def _seed_session(conn: sqlite3.Connection) -> tuple[int, int]:
    """Создает сессию с двумя связями: AUTO и HITL_PENDING.

    :return: идентификаторы связей (AUTO, HITL_PENDING).
    """
    db.create_session(
        conn,
        SESSION_ID,
        PROJECT_NAME,
        "rd.pdf",
        "smeta.xlsx",
        smeta_sheet="ЛС 02-01-03",
    )
    db.save_rd_items(
        conn,
        SESSION_ID,
        [
            {
                "row_idx": 1,
                "code": "РД-1",
                "description": "Кабель силовой ВВГнг 3х1.5",
                "unit": "м",
                "quantity": Decimal("100.0"),
            },
            {
                "row_idx": 2,
                "code": "РД-2",
                "description": "Арматура А500С d12",
                "unit": "т",
                "quantity": Decimal("10.0"),
            },
        ],
    )
    db.save_smeta_items(
        conn,
        SESSION_ID,
        [
            {
                "row_idx": 1,
                "code": "1.7-10-1",
                "description": "Кабель силовой ВВГнг 3х1.5",
                "unit": "м",
                "quantity": Decimal("100.0"),
                "unit_price": Decimal("2500.0"),
                "total_cost": Decimal("250000.0"),
            },
            {
                "row_idx": 2,
                "code": "4.8-175-2",
                "description": "Арматура А500С d12",
                "unit": "т",
                "quantity": Decimal("12.0"),
                "unit_price": Decimal("45000.0"),
                "total_cost": Decimal("540000.0"),
            },
        ],
    )

    rd_ids = [
        int(row["id"])
        for row in conn.execute(
            "SELECT id FROM rd_items WHERE session_id = ? ORDER BY id", (SESSION_ID,)
        )
    ]
    smeta_ids = [
        int(row["id"])
        for row in conn.execute(
            "SELECT id FROM smeta_items WHERE session_id = ? ORDER BY id", (SESSION_ID,)
        )
    ]

    db.save_matches(
        conn,
        SESSION_ID,
        [
            {
                "rd_item_id": rd_ids[0],
                "smeta_item_id": smeta_ids[0],
                "confidence_score": Decimal("95.0"),
                "conversion_ratio": Decimal("1.0"),
                "match_method": "EXACT",
                "status": "AUTO",
            },
            {
                "rd_item_id": rd_ids[1],
                "smeta_item_id": smeta_ids[1],
                "confidence_score": Decimal("60.0"),
                "conversion_ratio": Decimal("2.0"),
                "match_method": "FUZZY",
                "status": "HITL_PENDING",
            },
        ],
    )

    match_ids = [
        int(row["id"])
        for row in conn.execute(
            "SELECT id FROM match_records WHERE session_id = ? ORDER BY id", (SESSION_ID,)
        )
    ]
    return match_ids[0], match_ids[1]


def _seed_unmatched_smeta_session(
    conn: sqlite3.Connection,
    session_id: str,
    smeta_sheet: str,
) -> int:
    """Создает раздел, в котором позиция сметы не имеет пары в РД."""
    db.create_session(
        conn,
        session_id,
        PROJECT_NAME,
        f"{session_id}.pdf",
        "smeta.xlsx",
        smeta_sheet=smeta_sheet,
    )
    db.save_smeta_items(
        conn,
        session_id,
        [
            {
                "row_idx": 1,
                "code": "1.7-БЕЗ-РД",
                "description": "Позиция сметы без подтверждения в РД",
                "unit": "шт",
                "quantity": Decimal("10.0"),
                "unit_price": Decimal("0.0"),
                "total_cost": Decimal("0.0"),
            }
        ],
    )
    smeta_id = int(
        conn.execute(
            "SELECT id FROM smeta_items WHERE session_id = ?",
            (session_id,),
        ).fetchone()["id"]
    )
    db.save_matches(
        conn,
        session_id,
        [
            {
                "rd_item_id": None,
                "smeta_item_id": smeta_id,
                "confidence_score": Decimal("0.0"),
                "conversion_ratio": Decimal("1.0"),
                "match_method": "NO_MATCH",
                "status": "AUTO",
            }
        ],
    )
    match_id = int(
        conn.execute(
            "SELECT id FROM match_records WHERE session_id = ?",
            (session_id,),
        ).fetchone()["id"]
    )
    return match_id


def _open_sheet(path: Path):
    """Загружает книгу без вычисления формул и возвращает лист раздела."""
    workbook = load_workbook(path, data_only=False)
    return workbook.worksheets[1]


def _kpi_row(worksheet, label: str) -> int:
    """Находит номер KPI-строки по ее подписи на дашборде."""
    for row in range(
        exporter.DASHBOARD_FIRST_KPI_ROW,
        exporter.DASHBOARD_SECTION_TITLE_ROW,
    ):
        if worksheet[f"A{row}"].value == label:
            return row
    raise AssertionError(f"KPI-строка не найдена: {label}")


def test_exporter_creates_valid_xlsx(conn: sqlite3.Connection, tmp_path: Path) -> None:
    """Файл создается, содержит обязательный лист, шапку и строки данных."""
    _seed_session(conn)
    target = export_audit_to_excel(conn, SESSION_ID, tmp_path / "audit.xlsx")

    assert target.is_file()
    assert target.stat().st_size > 0

    workbook = load_workbook(target)
    assert workbook.sheetnames == [DASHBOARD_TITLE, "ЭОМ (ЛС 02-01-03)"]
    assert workbook.active.title == DASHBOARD_TITLE
    worksheet = workbook["ЭОМ (ЛС 02-01-03)"]

    assert worksheet["B1"].value == exporter.BACK_TO_DASHBOARD_FORMULA
    assert worksheet["B1"].font.color.rgb.endswith("FFFFFF")
    assert worksheet["B1"].font.bold is True
    assert worksheet["B1"].font.underline == "single"
    assert worksheet["C1"].value == PROJECT_NAME
    assert worksheet["B2"].value == "ID сессии:"
    assert worksheet["C2"].value == SESSION_ID
    assert worksheet["B3"].value == "Дата формирования:"
    assert worksheet["C3"].value
    assert worksheet["A4"].value == "ID"
    assert worksheet["C4"].value == "Позиция РД (наименование)"
    assert worksheet["Q4"].value == "Экспертный комментарий / Причина отклонения"
    assert worksheet.freeze_panes == "C5"
    assert worksheet.auto_filter.ref == "B4:Q6"
    assert worksheet.max_row == 7
    assert worksheet.max_column == 17
    assert not worksheet.column_dimensions["J"].hidden
    assert not worksheet.column_dimensions["K"].hidden
    assert not worksheet.column_dimensions["O"].hidden

    assert worksheet["B5"].value == 1
    assert worksheet["B6"].value == 2
    assert worksheet["E5"].number_format == exporter.VOLUME_FORMAT
    assert worksheet["J5"].number_format == exporter.MONEY_FORMAT
    assert worksheet["L5"].number_format == exporter.RATIO_FORMAT
    assert worksheet.column_dimensions["C"].width <= 50
    assert worksheet.column_dimensions["Q"].width <= 50

    # Нулевая дельта первой строки — мягкая зеленая заливка.
    assert worksheet["N5"].fill.fill_type == "solid"
    assert worksheet["N5"].fill.fgColor.rgb.endswith("C6EFCE")
    assert worksheet["P5"].value == "AUTO"
    assert worksheet["P6"].value == "HITL_PENDING"
    assert worksheet["P6"].fill.fgColor.rgb.endswith("FFEB9C")
    assert worksheet["P6"].font.color.rgb.endswith("9C6500")


def test_exporter_native_formulas_presence(
    conn: sqlite3.Connection, tmp_path: Path
) -> None:
    """Колонки M, N, O содержат формулы Excel, а не статичные числа."""
    _seed_session(conn)
    target = export_audit_to_excel(conn, SESSION_ID, tmp_path / "formulas.xlsx")
    worksheet = _open_sheet(target)

    for row in (5, 6):
        assert worksheet[f"K{row}"].value == f"=I{row}*J{row}"
        assert worksheet[f"M{row}"].value == f"=E{row}*L{row}"
        assert worksheet[f"N{row}"].value == f"=I{row}-M{row}"
        assert worksheet[f"O{row}"].value == f"=N{row}*J{row}"
        for column in ("M", "N", "O"):
            value = worksheet[f"{column}{row}"].value
            assert isinstance(value, str)
            assert value.startswith("=")
            assert worksheet[f"{column}{row}"].data_type == "f"


def test_exporter_hidden_id_column(conn: sqlite3.Connection, tmp_path: Path) -> None:
    """Служебная колонка A скрыта, но содержит идентификатор связи."""
    _, pending_match_id = _seed_session(conn)
    target = export_audit_to_excel(conn, SESSION_ID, tmp_path / "hidden.xlsx")
    worksheet = _open_sheet(target)

    assert worksheet.column_dimensions["A"].hidden is True
    assert worksheet["A4"].value == "ID"
    assert worksheet["A6"].value == pending_match_id


def test_exporter_adjusted_ratio_applied(
    conn: sqlite3.Connection, tmp_path: Path
) -> None:
    """Экспертная правка коэффициента попадает в колонку L вместо исходной."""
    _, pending_match_id = _seed_session(conn)

    # Две рецензии подряд: в отчет должна попасть последняя (1.25).
    db.record_expert_review(
        conn,
        pending_match_id,
        "ivanov",
        "ADJUSTED",
        comment="Первичное уточнение",
        adjusted_ratio=Decimal("1.5"),
    )
    db.record_expert_review(
        conn,
        pending_match_id,
        "petrov",
        "ADJUSTED",
        comment="Коэффициент уточнен по исполнительной схеме",
        adjusted_ratio=Decimal("1.25"),
    )

    target = export_audit_to_excel(conn, SESSION_ID, tmp_path / "adjusted.xlsx")
    worksheet = _open_sheet(target)

    assert worksheet["L6"].value == pytest.approx(1.25)
    assert worksheet["L6"].value != pytest.approx(2.0)
    assert worksheet["M6"].value == "=E6*L6"
    assert worksheet["P6"].value == "APPROVED"
    assert "исполнительной схеме" in str(worksheet["Q6"].value)

    stored_ratio = conn.execute(
        "SELECT conversion_ratio FROM match_records WHERE id = ?", (pending_match_id,)
    ).fetchone()["conversion_ratio"]
    assert stored_ratio == pytest.approx(2.0)


def test_exporter_rejected_reason_in_comment(
    conn: sqlite3.Connection, tmp_path: Path
) -> None:
    """REJECTED сохраняет наименование РД, причину в Q и не строит дельты."""
    _, pending_match_id = _seed_session(conn)
    db.record_expert_review(
        conn,
        pending_match_id,
        "ivanov",
        "REJECTED",
        reason_code="UNIT_MULTIPLICITY",
        comment="В смете расценка на 100 м",
    )

    target = export_audit_to_excel(conn, SESSION_ID, tmp_path / "rejected.xlsx")
    worksheet = _open_sheet(target)

    assert worksheet["P6"].value == "REJECTED"
    assert worksheet["C6"].value == "Арматура А500С d12"
    assert worksheet["E6"].value == pytest.approx(10.0)
    assert "UNIT_MULTIPLICITY" in str(worksheet["Q6"].value)
    assert "В смете расценка на 100 м" in str(worksheet["Q6"].value)
    # Отклоненная связь: дельтовые формулы не строятся.
    assert worksheet["N6"].value is None
    assert worksheet["O6"].value is None


def test_exporter_resolves_latest_session(conn: sqlite3.Connection, tmp_path: Path) -> None:
    """Сессия ``latest`` разрешается в последнюю созданную запись."""
    _seed_session(conn)
    db.create_session(conn, "export-session-002", "Другой объект", "rd2.pdf", "smeta2.xlsx")

    target = export_audit_to_excel(conn, "latest", tmp_path / "latest.xlsx")
    worksheet = _open_sheet(target)

    assert worksheet["C2"].value == "export-session-002"
    assert worksheet["C1"].value == "Другой объект"
    assert worksheet.max_row == 5

    with pytest.raises(LookupError):
        export_audit_to_excel(conn, "missing-session", tmp_path / "missing.xlsx")


def test_exporter_overrun_is_red(
    conn: sqlite3.Connection, tmp_path: Path
) -> None:
    """Превышение объема сметы над РД окрашивает статус и дельту в красный."""
    _seed_session(conn)
    conn.execute(
        """
        UPDATE smeta_items
        SET quantity = 150.0, total_cost = 375000.0
        WHERE session_id = ? AND row_idx = 1
        """,
        (SESSION_ID,),
    )
    conn.commit()

    target = export_audit_to_excel(conn, SESSION_ID, tmp_path / "overrun.xlsx")
    worksheet = _open_sheet(target)

    assert worksheet["P5"].value == "AUTO"
    assert worksheet["P5"].fill.fgColor.rgb.endswith("FFC7CE")
    assert worksheet["P5"].font.color.rgb.endswith("9C0006")
    assert worksheet["N5"].value == "=I5-M5"
    assert worksheet["N5"].fill.fgColor.rgb.endswith("FFC7CE")
    assert worksheet["O5"].fill.fgColor.rgb.endswith("FFC7CE")


def test_exporter_without_rd_pair_is_gray_and_has_no_deltas(
    conn: sqlite3.Connection, tmp_path: Path
) -> None:
    """Строка без пары РД не получает зеленый AUTO и дельтовые формулы."""
    match_id = _seed_unmatched_smeta_session(conn, "no-rd-session", "ЛС 02-01-01")

    target = export_audit_to_excel(conn, "no-rd-session", tmp_path / "no-rd.xlsx")
    worksheet = _open_sheet(target)

    assert worksheet["A5"].value == match_id
    assert worksheet["C5"].value is None
    assert worksheet["E5"].value is None
    assert worksheet["P5"].value == exporter.STATUS_NOT_IN_RD
    assert worksheet["P5"].fill.fgColor.rgb.endswith("D9D9D9")
    assert worksheet["P5"].font.color.rgb.endswith("595959")
    for column in ("L", "M", "N", "O"):
        assert worksheet[f"{column}5"].value is None
    assert worksheet["N5"].fill.fill_type is None

    # В разделе нет ни одной ненулевой цены: шумовые колонки скрываются.
    assert worksheet.column_dimensions["J"].hidden is True
    assert worksheet.column_dimensions["K"].hidden is True
    assert worksheet.column_dimensions["O"].hidden is True
    assert worksheet.auto_filter.ref == "B4:Q5"

    workbook = load_workbook(target)
    dashboard = workbook[DASHBOARD_TITLE]
    assert dashboard["A21"].value == "АР"
    assert dashboard["C21"].value == 1
    assert dashboard["D21"].value == 1
    assert dashboard["E21"].value == 0
    assert dashboard["F21"].value == 0
    assert dashboard["G21"].value == 0
    assert dashboard["J21"].value == "—"
    assert dashboard["K21"].value == "—"
    assert dashboard["L21"].value == exporter.VERDICT_NO_RD_SPEC
    assert dashboard["L21"].fill.fgColor.rgb.endswith("FFE699")


def test_export_sessions_readable_titles_and_dashboard(
    conn: sqlite3.Connection, tmp_path: Path
) -> None:
    """Многолистовая книга получает дашборд и вкладки «Раздел (Шифр)»."""
    _seed_session(conn)
    conn.execute(
        "UPDATE sessions SET smeta_sheet = 'ЛС 02-01-03' WHERE id = ?",
        (SESSION_ID,),
    )
    conn.commit()
    _seed_unmatched_smeta_session(conn, "ar-session", "ЛС 02-01-01")

    target = export_sessions_to_excel(
        conn,
        [SESSION_ID, "ar-session"],
        tmp_path / "multi.xlsx",
    )
    workbook = load_workbook(target)

    assert workbook.sheetnames == [
        DASHBOARD_TITLE,
        "ЭОМ (ЛС 02-01-03)",
        "АР (ЛС 02-01-01)",
    ]
    assert workbook.active.title == DASHBOARD_TITLE

    dashboard = workbook[DASHBOARD_TITLE]
    assert dashboard["A1"].value == "СВОДНЫЙ АУДИТ ОБЪЁМОВ"
    assert dashboard["B7"].value == 3
    assert dashboard["B8"].value == 2
    assert dashboard["B9"].value == pytest.approx(1 / 2)
    assert dashboard["B10"].value == 1
    assert dashboard["B11"].value == 1
    assert dashboard["B12"].value == 1
    assert dashboard["B13"].value == pytest.approx(790000.0)
    assert dashboard["B14"].value == pytest.approx(-360000.0)
    assert dashboard["A21"].value == "ЭОМ"
    assert dashboard["B21"].value == "ЛС 02-01-03"
    assert dashboard["C21"].value == 2
    assert dashboard["D21"].value == 1
    assert dashboard["E21"].value == 1
    assert dashboard["F21"].value == pytest.approx(1.0)
    assert dashboard["G21"].value == 1
    assert dashboard["L21"].value == exporter.VERDICT_DISCREPANCY
    assert dashboard["A22"].value == "АР"
    assert dashboard["B22"].value == "ЛС 02-01-01"
    assert dashboard["C22"].value == 1
    assert dashboard["D22"].value == 1
    assert dashboard["E22"].value == 0
    assert dashboard["F22"].value == 0
    assert dashboard["G22"].value == 0
    assert dashboard["J22"].value == "—"
    assert dashboard["K22"].value == "—"
    assert dashboard["L22"].value == exporter.VERDICT_NO_RD_SPEC
    assert dashboard["L22"].fill.fgColor.rgb.endswith("FFE699")
    assert dashboard["A21"].hyperlink is not None
    assert dashboard.auto_filter.ref == "A20:L22"

    eom = workbook["ЭОМ (ЛС 02-01-03)"]
    assert eom.auto_filter.ref == "B4:Q6"
    assert eom.freeze_panes == "C5"

    ar = workbook["АР (ЛС 02-01-01)"]
    assert ar["P5"].value == exporter.STATUS_NOT_IN_RD
    assert ar["C5"].value is None
    assert ar["N5"].value is None
    assert ar["P5"].fill.fgColor.rgb.endswith("D9D9D9")


def test_dashboard_percentage_format(
    conn: sqlite3.Connection, tmp_path: Path
) -> None:
    """Покрытие выводится нативным процентом в KPI и таблице разделов."""
    _seed_session(conn)
    target = export_audit_to_excel(conn, SESSION_ID, tmp_path / "percent.xlsx")
    dashboard = load_workbook(target)[DASHBOARD_TITLE]

    kpi_row = _kpi_row(dashboard, "Покрытие спецификациями РД")
    assert dashboard[f"B{kpi_row}"].number_format == exporter.PERCENT_FORMAT
    assert dashboard["F21"].number_format == exporter.PERCENT_FORMAT


def test_dashboard_dash_for_missing_prices(
    conn: sqlite3.Connection, tmp_path: Path
) -> None:
    """Без подключенных цен дашборд показывает центрированный прочерк."""
    _seed_unmatched_smeta_session(conn, "dash-session", "ЛС 02-01-01")
    target = export_audit_to_excel(conn, "dash-session", tmp_path / "dash.xlsx")
    dashboard = load_workbook(target)[DASHBOARD_TITLE]

    cost_row = _kpi_row(dashboard, "Стоимость сметы, руб.")
    assert dashboard[f"B{cost_row}"].value == "—"
    assert dashboard[f"B{cost_row}"].alignment.horizontal == "center"
    assert dashboard["J21"].value == "—"
    assert dashboard["K21"].value == "—"
    assert dashboard["J21"].alignment.horizontal == "center"
    assert dashboard["K21"].alignment.horizontal == "center"


def test_sheet_has_summary_total_row(
    conn: sqlite3.Connection, tmp_path: Path
) -> None:
    """Внизу листа раздела есть бухгалтерская строка ИТОГО с формулами SUM."""
    _seed_session(conn)
    target = export_audit_to_excel(conn, SESSION_ID, tmp_path / "totals.xlsx")
    worksheet = _open_sheet(target)

    summary_row = exporter.FIRST_DATA_ROW + 2
    assert worksheet[f"B{summary_row}"].value == "ИТОГО ПО РАЗДЕЛУ:"
    for column in ("E", "I", "M", "N"):
        assert (
            worksheet[f"{column}{summary_row}"].value
            == f"=SUM({column}5:{column}6)"
        )
        assert worksheet[f"{column}{summary_row}"].number_format == (
            exporter.VOLUME_FORMAT
        )
    assert worksheet[f"B{summary_row}"].font.bold is True
    assert worksheet[f"N{summary_row}"].font.bold is True
    assert worksheet[f"E{summary_row}"].border.top.style == "thin"
    assert worksheet[f"E{summary_row}"].border.bottom.style == "double"


def test_return_link_styling(conn: sqlite3.Connection, tmp_path: Path) -> None:
    """Плашка возврата на дашборд имеет темно-синюю заливку и белый шрифт."""
    _seed_session(conn)
    target = export_audit_to_excel(conn, SESSION_ID, tmp_path / "link.xlsx")
    worksheet = _open_sheet(target)

    assert worksheet["B1"].value == exporter.BACK_TO_DASHBOARD_FORMULA
    assert worksheet["B1"].fill.fgColor.rgb.endswith("1F4E79")
    assert worksheet["B1"].font.color.rgb.endswith("FFFFFF")
    assert worksheet["B1"].font.bold is True
    assert worksheet["B1"].alignment.horizontal == "center"
    assert worksheet.column_dimensions["B"].width >= 20


def test_readable_sheet_title_respects_excel_limits(tmp_path: Path) -> None:
    """Имя вкладки содержит читаемый раздел и очищается от запрещенных знаков."""
    assert (
        exporter._readable_sheet_title("ЛС 02-01-16", "fallback")
        == "СДС (ЛС 02-01-16)"
    )

    raw = "Очень длинное название / раздела ? * : [ ] (ЛС 02-01-99)"
    title = exporter._readable_sheet_title(raw, "fallback")
    assert len(title) <= 31
    for invalid in "\\/?*:[]":
        assert invalid not in title

