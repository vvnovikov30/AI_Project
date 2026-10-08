import os
import re
from decimal import Decimal
import openpyxl
import pytest

from models import UnitClass
from parsers import LSRParser, PDFSpecParser


# ==============================================================================
# 1. МОДУЛЬНЫЕ ТЕСТЫ LSRParser (СИНТЕТИЧЕСКИЕ ДАННЫЕ В ПАМЯТИ)
# ==============================================================================

class TestLSRParserSynthetic:
    """Проверка логики парсера смет без зависимости от внешних файлов."""

    def test_clean_estimate_title(self):
        """Очистка наименований от формул ценообразования, префиксов и артефактов."""
        raw_1 = (
            "ОБОРУДОВАНИЕ:\n"
            "Цилиндрическая IP-видеокамера с объективом 4 Мп\n"
            "1 236,60 = [10 913 / 1,22 / 7,54] + 3% Трансп + 1,2% Заг.скл"
        )
        cleaned_1 = LSRParser._clean_estimate_title(raw_1)
        assert cleaned_1 == "цилиндрическая ip-видеокамера с объективом 4 мп"

        raw_2 = "ОБОРУДОВАНИЕ:_x000D_\nКлавиатура компьютерная USB."
        cleaned_2 = LSRParser._clean_estimate_title(raw_2)
        assert cleaned_2 == "клавиатура компьютерная usb."

        raw_3 = "Монтаж наружной камеры видеонаблюдения"
        cleaned_3 = LSRParser._clean_estimate_title(raw_3)
        assert cleaned_3 == "монтаж наружной камеры видеонаблюдения"

    def test_merged_cells_resolution(self, tmp_path):
        """Корректное извлечение значений из объединенных ячеек Excel."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="Объединенный текст")
        ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=2)

        file_path = tmp_path / "test_merged.xlsx"
        wb.save(file_path)

        parser = LSRParser(str(file_path))
        merge_map = parser._resolve_merged_cells_map(parser.wb.active)

        assert merge_map[(1, 1)] == "Объединенный текст"
        assert merge_map[(1, 2)] == "Объединенный текст"
        assert merge_map[(2, 1)] == "Объединенный текст"
        assert merge_map[(2, 2)] == "Объединенный текст"

    def test_noise_filtering_and_hierarchy(self, tmp_path):
        """Фильтрация служебных строк (НР, СП, ЗТР, шапки) и построение дерева parent_id."""
        wb = openpyxl.Workbook()
        ws = wb.active

        # Шапка и раздел
        ws.cell(row=1, column=1, value="ЛОКАЛЬНАЯ СМЕТА № 02-01-14")
        ws.cell(row=2, column=1, value="Раздел: Оборудование")

        # 1. Базовая расценка на монтаж (№ 1)
        ws.cell(row=3, column=1, value="1")
        ws.cell(row=3, column=2, value="4.10-174-2\nПоправка: Гл.12")
        ws.cell(row=3, column=3, value="Монтаж наружной камеры видеонаблюдения")
        ws.cell(row=3, column=4, value="10 шт.")
        ws.cell(row=3, column=5, value="0,5")

        # Служебный шум (должен быть проигнорирован)
        ws.cell(row=4, column=1, value=None)
        ws.cell(row=4, column=3, value="НР от ФОТ")
        ws.cell(row=4, column=4, value="%")
        ws.cell(row=4, column=5, value=90)

        ws.cell(row=5, column=1, value=None)
        ws.cell(row=5, column=3, value="СП от ФОТ")
        ws.cell(row=5, column=4, value="%")
        ws.cell(row=5, column=5, value=46)

        ws.cell(row=6, column=1, value=None)
        ws.cell(row=6, column=3, value="ЗТР")
        ws.cell(row=6, column=4, value="чел-ч")
        ws.cell(row=6, column=5, value=25.33)

        # 2. Дочерний ресурс оборудования (№ 2)
        ws.cell(row=7, column=1, value="2")
        ws.cell(row=7, column=2, value="Прайс-лист")
        ws.cell(row=7, column=3, value="ОБОРУДОВАНИЕ:\nКамера уличная 4Мп")
        ws.cell(row=7, column=4, value="ШТ")
        ws.cell(row=7, column=5, value=5)

        # 3. Служебная строка-нумерация граф таблицы: 1 | 2 | 3 | 4 | 5
        ws.cell(row=8, column=1, value="1")
        ws.cell(row=8, column=2, value="2")
        ws.cell(row=8, column=3, value="3")
        ws.cell(row=8, column=4, value="4")
        ws.cell(row=8, column=5, value=5)

        # 4. Новый раздел и расценка
        ws.cell(row=9, column=1, value="Раздел: Кабельные изделия")
        ws.cell(row=10, column=1, value="3")
        ws.cell(row=10, column=2, value="4.8-280-2")
        ws.cell(row=10, column=3, value="Прокладка пластикового кабель-канала")
        ws.cell(row=10, column=4, value="100 м")
        ws.cell(row=10, column=5, value=3)

        # 5. Материал внутри расценки с дробным номером (№ 3,1)
        ws.cell(row=11, column=1, value="3,1")
        ws.cell(row=11, column=2, value="1.1-1-1962")
        ws.cell(row=11, column=3, value="Кабель-канал 10х15 мм")
        ws.cell(row=11, column=4, value="м")
        ws.cell(row=11, column=5, value=300)

        file_path = tmp_path / "test_hierarchy.xlsx"
        wb.save(file_path)

        parser = LSRParser(str(file_path))
        items = parser.parse()

        # Ровно 4 валидные позиции (1, 2, 3, 3,1)
        assert len(items) == 4

        # Проверка базовой расценки № 1
        it_work = items[0]
        assert it_work.code == "4.10-174-2"
        assert not it_work.is_resource
        assert it_work.parent_id is None
        assert it_work.volume == Decimal("5.0")  # 0.5 * 10
        assert it_work.canonical_unit == UnitClass.COUNT
        assert it_work.context_section == "Раздел: Оборудование"

        # Проверка дочернего ресурса № 2
        it_res = items[1]
        assert it_res.code == "Прайс-лист"
        assert it_res.is_resource
        assert it_res.parent_id == "LSR_3"  # Привязан к родительской расценке строки 3
        assert it_res.volume == Decimal("5")
        assert it_res.clean_title == "камера уличная 4мп"

        # Проверка сброса parent_id при новой расценке
        it_work2 = items[2]
        assert it_work2.code == "4.8-280-2"
        assert not it_work2.is_resource
        assert it_work2.parent_id is None
        assert it_work2.volume == Decimal("300.0")  # 3 * 100
        assert it_work2.canonical_unit == UnitClass.LENGTH
        assert it_work2.context_section == "Раздел: Кабельные изделия"

        # Проверка дочернего материала № 3,1
        it_res2 = items[3]
        assert it_res2.code == "1.1-1-1962"
        assert it_res2.is_resource
        assert it_res2.parent_id == "LSR_10"
        assert it_res2.volume == Decimal("300")


# ==============================================================================
# 2. МОДУЛЬНЫЕ ТЕСТЫ PDFSpecParser
# ==============================================================================

class TestPDFSpecParserUnit:
    """Проверка вспомогательных методов извлечения спецификации."""

    def test_crop_box_calculation(self):
        """Проверка формулы отсечения штампа по ГОСТ 21.110 (55 мм снизу)."""
        class DummyPage:
            width = 842.0   # A4 альбомная
            height = 595.0

        parser = PDFSpecParser("dummy.pdf", debug=False)
        box = parser._get_crop_box(DummyPage())

        stamp_h = 55 * 2.83465
        assert box[0] == 0
        assert box[1] == 0
        assert box[2] == 842.0
        assert pytest.approx(box[3], 0.01) == 595.0 - stamp_h


# ==============================================================================
# 3. ИНТЕГРАЦИОННЫЕ ТЕСТЫ НА РЕАЛЬНЫХ ДАННЫХ ПРОЕКТА
# ==============================================================================

LSR_FILE = "data/сот.xlsx"
SPEC_FILE = "data/сот2.pdf"

@pytest.mark.skipif(not os.path.exists(LSR_FILE), reason="Файл сметы сот.xlsx не найден")
def test_real_lsr_file_clean_extraction():
    """Интеграционный тест: реальная смета ТСН-2001 содержит ровно 44 чистые позиции."""
    parser = LSRParser(LSR_FILE)
    items = parser.parse()

    assert len(items) == 44

    # Проверка ключевых позиций
    camera_work = next(it for it in items if it.code == "4.10-174-2")
    assert camera_work.volume == Decimal("5.0")
    assert not camera_work.is_resource

    cable_item = next(it for it in items if "витая пара" in it.clean_title)
    assert cable_item.volume == Decimal("1326")
    assert cable_item.canonical_unit == UnitClass.LENGTH


@pytest.mark.skipif(not os.path.exists(SPEC_FILE), reason="Файл спецификации сот2.pdf не найден")
def test_real_pdf_spec_file_clean_extraction():
    """Интеграционный тест: реальная спецификация ГОСТ 21.110 содержит ровно 23 позиции."""
    parser = PDFSpecParser(SPEC_FILE, debug=False)
    items = parser.extract_items()

    assert len(items) == 23

    # Проверка отсутствия фантомных строк с номерами колонок таблицы
    assert not any(it.raw_title.strip() == "2" for it in items)
    assert not any("оборудование:" in it.clean_title and len(it.clean_title) < 18 for it in items)

    # Проверка извлечения ключевых проектных позиций
    pipe_item = next(it for it in items if "гофротруба" in it.clean_title)
    assert pipe_item.volume == Decimal("1000")
    assert pipe_item.canonical_unit == UnitClass.LENGTH