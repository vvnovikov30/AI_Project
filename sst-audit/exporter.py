"""Генератор ведомости аудита объемов в формате Excel для контролирующих органов.

Отчет строится по данным SQLite (SSOT): строки сшиваются через представление
``v_audit_export``, при наличии нескольких экспертных рецензий выбирается
последняя. Расчетные колонки (``M``, ``N``, ``O``) записываются нативными
формулами Excel, поэтому проверяющий видит не только результат, но и логику
расчета и может пересчитать ведомость средствами самого Excel.

Книга начинается с титульного листа «📊 Сводный дашборд» с KPI объекта.
Вкладки разделов получают читаемые имена вида ``АР (ЛС 02-01-01)``. Зеленая
заливка статуса допустима только при наличии реальной пары РД и дельте
объема в пределах допуска; несопоставленные строки нейтрально-серые.
Покрытие спецификациями рассчитывается отдельно по оборудованию и
материалам (главы 1, 2, 13 и «Прайс-лист»), а монтажные и строительные
работы (главы 3, 4, 6) показываются как норматив СМР. На каждом листе
раздела под таблицей выводится бухгалтерская строка «ИТОГО ПО РАЗДЕЛУ».

Запуск как самостоятельного скрипта::

    python exporter.py --db sst_audit.db --session latest --out report.xlsx
"""

from __future__ import annotations

import argparse
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import sqlite3
from typing import Any, Final, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

import db
from evaluator import DiscrepancyDetector

SHEET_TITLE: Final[str] = "Ведомость аудита объемов"
DASHBOARD_TITLE: Final[str] = "📊 Сводный дашборд"
BACK_TO_DASHBOARD_FORMULA: Final[str] = (
    "=HYPERLINK(\"#'" + DASHBOARD_TITLE + "'!A1\", \"⬅ Вернуться на дашборд\")"
)
LATEST_SESSION: Final[str] = "latest"
DEFAULT_OUTPUT: Final[str] = "audit_report.xlsx"

HEADER_ROW: Final[int] = 4
FIRST_DATA_ROW: Final[int] = 5
HIDDEN_ID_COLUMN: Final[str] = "A"
ONE: Final[Decimal] = Decimal("1.0")
ZERO: Final[Decimal] = Decimal("0")

MONEY_FORMAT: Final[str] = '#,##0.00 "₽"'
VOLUME_FORMAT: Final[str] = "#,##0.000"
RATIO_FORMAT: Final[str] = "#,##0.0000"
PERCENT_FORMAT: Final[str] = "0.0%"
INTEGER_FORMAT: Final[str] = "#,##0"

# Классификация сметных строк ТСН-2001. Оборудование и материалы подлежат
# спецификации в РД; монтажные и строительные работы являются нормативом СМР
# и не имеют отдельных строк в таблицах спецификаций.
KIND_EQUIPMENT: Final[str] = "EQUIPMENT"
KIND_SMR: Final[str] = "SMR"
KIND_UNCLASSIFIED: Final[str] = "UNCLASSIFIED"
EQUIPMENT_CHAPTERS: Final[frozenset[str]] = frozenset({"1", "2", "13"})
SMR_CHAPTERS: Final[frozenset[str]] = frozenset({"3", "4", "6"})
_RATE_CHAPTER_PATTERN: Final[re.Pattern[str]] = re.compile(r"^\s*(\d{1,2})[.\-]")

STATUS_AUTO: Final[str] = "AUTO"
STATUS_APPROVED: Final[str] = "APPROVED"
STATUS_REJECTED: Final[str] = "REJECTED"
STATUS_HITL_PENDING: Final[str] = "HITL_PENDING"
STATUS_NOT_IN_RD: Final[str] = "НЕТ В РД"
STATUS_NOT_IN_SMETA: Final[str] = "НЕТ В СМЕТЕ"

# Читаемые аббревиатуры разделов. Шифр извлекается из имени листа сметы
# (``ЛС 02-01-03``), поэтому название вкладки не зависит от порядка разделов.
SECTION_NAMES_BY_CODE: Final[dict[str, str]] = {
    "02-01-01": "АР",
    "02-01-03": "ЭОМ",
    "02-01-04": "ВК",
    "02-01-05": "ВК",
    "02-01-06": "ОВ",
    "02-01-08": "АПС",
    "02-01-09": "СОУЭ",
    "02-01-11": "РТС",
    "02-01-12": "СКТ",
    "02-01-13": "СКС",
    "02-01-14": "СОТ",
    "02-01-15": "СКУД",
    "02-01-16": "СДС",
}

VERDICT_OK: Final[str] = "ПРОВЕРЕНО"
VERDICT_DISCREPANCY: Final[str] = "РАСХОЖДЕНИЯ"
VERDICT_REVIEW: Final[str] = "ТРЕБУЕТ ПРОВЕРКИ"
VERDICT_NO_MATCH: Final[str] = "⚠️ Позиции РД не сопоставлены со сметой"
VERDICT_NO_RD_SPEC: Final[str] = (
    "⚠️ В РД не найдены таблицы спецификаций (проверьте PDF)"
)
VERDICT_NO_DATA: Final[str] = "НЕТ ДАННЫХ"

DASHBOARD_KPI_HEADER_ROW: Final[int] = 6
DASHBOARD_FIRST_KPI_ROW: Final[int] = 7
DASHBOARD_SECTION_TITLE_ROW: Final[int] = 19
DASHBOARD_SECTION_HEADER_ROW: Final[int] = 20
DASHBOARD_FIRST_SECTION_ROW: Final[int] = 21

REJECT_REASON_TITLES: Final[dict[str, str]] = {
    "SEMANTIC_MISMATCH": "Семантика: описание относится к другому ресурсу",
    "UNIT_MULTIPLICITY": "Единицы измерения или кратность (например, 100 м)",
    "PARSER_ERROR": "Ошибка парсера: данные извлечены неверно",
    "OTHER": "Иное (см. комментарий эксперта)",
}

COLUMN_TITLES: Final[dict[str, str]] = {
    "A": "ID",
    "B": "№ п/п",
    "C": "Позиция РД (наименование)",
    "D": "Ед. изм. РД",
    "E": "Объем по РД",
    "F": "Шифр расценки (Смета)",
    "G": "Наименование расценки (Смета)",
    "H": "Ед. изм. Сметы",
    "I": "Объем по Смете",
    "J": "Цена за ед., руб.",
    "K": "Стоимость по Смете, руб.",
    "L": "Коэфф. приведения",
    "M": "Объем РД приведенный",
    "N": "Дельта объема",
    "O": "Дельта стоимости, руб.",
    "P": "Статус аудита",
    "Q": "Экспертный комментарий / Причина отклонения",
}

# Ограничения автоподбора ширины: длинные текстовые колонки не должны
# вытеснять числовые данные за пределы экрана.
COLUMN_WIDTH_LIMITS: Final[dict[str, int]] = {
    "A": 12,
    "B": 7,
    "C": 50,
    "D": 10,
    "E": 14,
    "F": 24,
    "G": 50,
    "H": 10,
    "I": 14,
    "J": 15,
    "K": 19,
    "L": 14,
    "M": 18,
    "N": 14,
    "O": 19,
    "P": 16,
    "Q": 50,
}

AUDIT_ROWS_SQL: Final[str] = """
SELECT v.*
FROM v_audit_export AS v
WHERE v.session_id = ?
  AND (
      v.expert_review_id IS NULL
      OR v.expert_review_id = (
          SELECT MAX(er.id)
          FROM expert_reviews AS er
          WHERE er.match_record_id = v.match_id
      )
  )
ORDER BY v.match_id
"""


def classify_smeta_code(code: Optional[str]) -> str:
    """Классифицирует сметную строку как оборудование, СМР или прочее.

    В ТСН-2001 шифр расценки начинается с номера главы: главы 1, 2 и 13
    относятся к оборудованию и материалам, главы 3, 4 и 6 — к монтажным и
    строительным работам. Строки с пометкой «Прайс-лист» или «Оборудование»
    также считаются оборудованием.
    """
    normalized = (code or "").strip().casefold()
    if not normalized:
        return KIND_UNCLASSIFIED
    if "прайс" in normalized or "оборуд" in normalized:
        return KIND_EQUIPMENT

    match = _RATE_CHAPTER_PATTERN.match(normalized)
    if match is None:
        return KIND_UNCLASSIFIED
    chapter = match.group(1)
    if chapter in EQUIPMENT_CHAPTERS:
        return KIND_EQUIPMENT
    if chapter in SMR_CHAPTERS:
        return KIND_SMR
    return KIND_UNCLASSIFIED


@dataclass(frozen=True)
class AuditRecord:
    """Нормализованная строка ведомости, готовая к записи в Excel."""

    match_id: int
    rd_item_id: Optional[int]
    smeta_item_id: Optional[int]
    rd_description: Optional[str]
    rd_unit: Optional[str]
    rd_quantity: Optional[Decimal]
    smeta_code: Optional[str]
    smeta_description: Optional[str]
    smeta_unit: Optional[str]
    smeta_quantity: Optional[Decimal]
    unit_price: Optional[Decimal]
    total_cost: Optional[Decimal]
    ratio: Decimal
    status: str
    comment: Optional[str]

    @property
    def has_rd_pair(self) -> bool:
        """Есть ли реальная позиция РД, а не только запись связи."""
        return self.rd_item_id is not None and self.rd_description is not None

    @property
    def has_smeta_item(self) -> bool:
        """Есть ли в связи позиция сметы."""
        return self.smeta_item_id is not None and self.smeta_description is not None

    @property
    def has_pair(self) -> bool:
        """Есть ли полноценная пара «смета ↔ РД» для сравнения объемов."""
        return self.has_rd_pair and self.has_smeta_item

    @property
    def smeta_kind(self) -> str:
        """Тип сметной строки: оборудование, СМР или вне классификации."""
        return classify_smeta_code(self.smeta_code)

    @property
    def adjusted_rd_volume(self) -> Optional[Decimal]:
        """Приведенный объем РД: E * L (без нормативов отходов)."""
        if self.rd_quantity is None:
            return None
        return self.rd_quantity * self.ratio

    @property
    def delta_volume(self) -> Optional[Decimal]:
        """Дельта объема: смета минус приведенный объем РД."""
        adjusted = self.adjusted_rd_volume
        if adjusted is None or self.smeta_quantity is None:
            return None
        return self.smeta_quantity - adjusted

    @property
    def estimated_cost(self) -> Optional[Decimal]:
        """Сметная стоимость строки: ``total_cost`` либо ``I * J``.

        Если парсер сохранил только цену и объем, стоимость выводится
        расчетным путем — так финансовая реконсиляция не зависит от наличия
        предрасчитанного поля в смете.
        """
        if self.total_cost is not None:
            return self.total_cost
        if self.smeta_quantity is not None and self.unit_price is not None:
            return self.smeta_quantity * self.unit_price
        return None

    @property
    def tolerance(self) -> Optional[Decimal]:
        """Допустимое расхождение по регламенту ``DiscrepancyDetector``."""
        adjusted = self.adjusted_rd_volume
        if adjusted is None:
            return None
        return (
            DiscrepancyDetector.ABS_TOLERANCE
            + DiscrepancyDetector.REL_TOLERANCE * adjusted
        )

    @property
    def is_within_tolerance(self) -> bool:
        """Признак совпадения объемов в пределах допуска."""
        delta = self.delta_volume
        tolerance = self.tolerance
        if delta is None or tolerance is None:
            return False
        return abs(delta) <= tolerance

    @property
    def is_overrun(self) -> bool:
        """Признак превышения сметы относительно приведенного объема РД."""
        delta = self.delta_volume
        tolerance = self.tolerance
        if delta is None or tolerance is None:
            return False
        return delta > tolerance


# ----------------------------------------------------------------------
# 1. Нормализация значений SQLite -> Excel
# ----------------------------------------------------------------------


def _to_decimal(value: Any) -> Optional[Decimal]:
    """Приводит значение SQLite к Decimal без потери десятичного представления."""
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None


def _to_excel_number(value: Any) -> Optional[float]:
    """Возвращает число для записи в ячейку openpyxl.

    openpyxl не поддерживает ``Decimal``, поэтому конвертация выполняется
    только на границе формирования отчета; все сравнения и допуски считаются
    в Decimal.
    """
    number = _to_decimal(value)
    return None if number is None else float(number)


def _to_text(value: Any) -> Optional[str]:
    """Возвращает непустую строку либо ``None``."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _to_optional_int(value: Any) -> Optional[int]:
    """Приводит SQLite-идентификатор к int, сохраняя ``None``."""
    if value is None:
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


# ----------------------------------------------------------------------
# 2. Чтение SSOT и нормализация строк отчета
# ----------------------------------------------------------------------


def resolve_session_id(conn: sqlite3.Connection, session_id: str) -> str:
    """Разрешает ``latest`` в идентификатор последней созданной сессии.

    :raises LookupError: в базе нет сессий.
    """
    if session_id != LATEST_SESSION:
        return session_id

    row = conn.execute(
        "SELECT id FROM sessions ORDER BY created_at DESC, rowid DESC LIMIT 1"
    ).fetchone()
    if row is None:
        raise LookupError("В базе данных нет ни одной сессии аудита")
    return str(row["id"])


def _load_session(conn: sqlite3.Connection, session_id: str) -> sqlite3.Row:
    """Возвращает метаданные сессии.

    :raises LookupError: сессия с указанным идентификатором не найдена.
    """
    row = conn.execute(
        """
        SELECT id, project_name, rd_filename, smeta_filename, smeta_sheet,
               created_at, status
        FROM sessions
        WHERE id = ?
        """,
        (session_id,),
    ).fetchone()
    if row is None:
        raise LookupError(f"Сессия '{session_id}' не найдена")
    return row


def _effective_ratio(row: sqlite3.Row) -> Decimal:
    """Экспертный коэффициент приоритетнее исходного коэффициента матчинга."""
    adjusted = _to_decimal(row["adjusted_ratio"])
    if adjusted is not None and adjusted > 0:
        return adjusted
    ratio = _to_decimal(row["conversion_ratio"])
    if ratio is not None and ratio > 0:
        return ratio
    return ONE


def _audit_status(row: sqlite3.Row) -> str:
    """Сводит статус связи и решение эксперта к статусу ведомости.

    Отсутствие позиции РД имеет приоритет над техническим ``AUTO``: иначе
    несопоставленная строка выглядит как успешно проверенная. ``ADJUSTED``
    показывается как ``APPROVED``: эксперт принял связь после уточнения
    коэффициента. Аномальный ``RESOLVED`` без решения эксперта возвращается
    в очередь, чтобы не маскировать нарушение пайплайна.
    """
    match_status = _to_text(row["match_status"]) or STATUS_HITL_PENDING
    decision = _to_text(row["decision"])
    has_rd = (
        _to_optional_int(row["rd_item_id"]) is not None
        and _to_text(row["rd_description"]) is not None
    )
    has_smeta = (
        _to_optional_int(row["smeta_item_id"]) is not None
        and _to_text(row["smeta_description"]) is not None
    )

    if decision == "REJECTED":
        return STATUS_REJECTED
    if not has_rd:
        return STATUS_NOT_IN_RD
    if not has_smeta:
        return STATUS_NOT_IN_SMETA
    if decision in {"APPROVED", "ADJUSTED"}:
        return STATUS_APPROVED
    if match_status == STATUS_AUTO:
        return STATUS_AUTO
    return STATUS_HITL_PENDING


def _compose_comment(row: sqlite3.Row) -> Optional[str]:
    """Склеивает причину отклонения и экспертный комментарий для колонки Q."""
    parts: list[str] = []
    reason_code = _to_text(row["reject_reason_code"])
    if reason_code:
        reason_title = REJECT_REASON_TITLES.get(reason_code, reason_code)
        parts.append(f"{reason_code}: {reason_title}")

    comment = _to_text(row["review_comment"])
    if comment:
        parts.append(comment)
    return ". ".join(parts) if parts else None


def _record_from_row(row: sqlite3.Row) -> AuditRecord:
    """Преобразует строку ``v_audit_export`` в модель ведомости."""
    return AuditRecord(
        match_id=int(row["match_id"]),
        rd_item_id=_to_optional_int(row["rd_item_id"]),
        smeta_item_id=_to_optional_int(row["smeta_item_id"]),
        rd_description=_to_text(row["rd_description"]),
        rd_unit=_to_text(row["rd_unit"]),
        rd_quantity=_to_decimal(row["rd_quantity"]),
        smeta_code=_to_text(row["smeta_code"]),
        smeta_description=_to_text(row["smeta_description"]),
        smeta_unit=_to_text(row["smeta_unit"]),
        smeta_quantity=_to_decimal(row["smeta_quantity"]),
        unit_price=_to_decimal(row["unit_price"]),
        total_cost=_to_decimal(row["total_cost"]),
        ratio=_effective_ratio(row),
        status=_audit_status(row),
        comment=_compose_comment(row),
    )


def load_audit_records(
    conn: sqlite3.Connection,
    session_id: str,
) -> list[AuditRecord]:
    """Возвращает строки ведомости по сессии (по одной на связь)."""
    rows = conn.execute(AUDIT_ROWS_SQL, (session_id,)).fetchall()
    return [_record_from_row(row) for row in rows]


# ----------------------------------------------------------------------
# 3. Стили и оформление листа
# ----------------------------------------------------------------------

_HEADER_FONT: Final[Font] = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
_HEADER_FILL: Final[PatternFill] = PatternFill("solid", fgColor="1F4E78")
_HEADER_ALIGNMENT: Final[Alignment] = Alignment(
    horizontal="center", vertical="center", wrap_text=True
)
_META_LABEL_FONT: Final[Font] = Font(name="Calibri", size=11, bold=True)
_META_VALUE_FONT: Final[Font] = Font(name="Calibri", size=11)
_PROJECT_FONT: Final[Font] = Font(name="Calibri", size=14, bold=True, color="1F4E78")
_BACK_LINK_FONT: Final[Font] = Font(
    name="Calibri", size=9, bold=True, color="FFFFFF", underline="single"
)
_BACK_LINK_FILL: Final[PatternFill] = PatternFill(
    "solid", fgColor="1F4E79"
)
_BODY_FONT: Final[Font] = Font(name="Calibri", size=10)
_SUMMARY_FONT: Final[Font] = Font(name="Calibri", size=10, bold=True)
_SUMMARY_FILL: Final[PatternFill] = PatternFill("solid", fgColor="EAF1F8")
_THIN_SIDE: Final[Side] = Side(style="thin", color="B7C3D0")
_SUMMARY_TOP_SIDE: Final[Side] = Side(style="thin", color="1F4E79")
_SUMMARY_BOTTOM_SIDE: Final[Side] = Side(style="double", color="1F4E79")
_CELL_BORDER: Final[Border] = Border(
    left=_THIN_SIDE,
    right=_THIN_SIDE,
    top=_THIN_SIDE,
    bottom=_THIN_SIDE,
)
_SUMMARY_BORDER: Final[Border] = Border(
    left=_THIN_SIDE,
    right=_THIN_SIDE,
    top=_SUMMARY_TOP_SIDE,
    bottom=_SUMMARY_BOTTOM_SIDE,
)
_GREEN_FILL: Final[PatternFill] = PatternFill("solid", fgColor="C6EFCE")
_RED_FILL: Final[PatternFill] = PatternFill("solid", fgColor="FFC7CE")
_YELLOW_FILL: Final[PatternFill] = PatternFill("solid", fgColor="FFEB9C")
_WARNING_FILL: Final[PatternFill] = PatternFill("solid", fgColor="FFE699")
_GRAY_FILL: Final[PatternFill] = PatternFill("solid", fgColor="D9D9D9")
_WHITE_FILL: Final[PatternFill] = PatternFill("solid", fgColor="FFFFFF")

_GREEN_FONT: Final[Font] = Font(
    name="Calibri", size=10, bold=True, color="006100"
)
_RED_FONT: Final[Font] = Font(
    name="Calibri", size=10, bold=True, color="9C0006"
)
_YELLOW_FONT: Final[Font] = Font(
    name="Calibri", size=10, bold=True, color="9C6500"
)
_WARNING_FONT: Final[Font] = Font(
    name="Calibri", size=10, bold=True, color="9C5700"
)
_GRAY_FONT: Final[Font] = Font(
    name="Calibri", size=10, bold=True, color="595959"
)

_NO_PAIR_STATUSES: Final[frozenset[str]] = frozenset(
    {STATUS_NOT_IN_RD, STATUS_NOT_IN_SMETA}
)


def _record_status_style(record: AuditRecord) -> tuple[PatternFill, Font]:
    """Возвращает заливку и шрифт статуса без «зеленой лжи».

    Зеленый цвет допустим только для реальной пары РД/смета и дельты
    в пределах допуска. Несопоставленные строки всегда нейтрально-серые,
    HITL — желтый, отклоненные и превышения — красные.
    """
    if record.status in _NO_PAIR_STATUSES or not record.has_rd_pair:
        return _GRAY_FILL, _GRAY_FONT
    if record.status == STATUS_HITL_PENDING:
        return _YELLOW_FILL, _YELLOW_FONT
    if record.status == STATUS_REJECTED:
        return _RED_FILL, _RED_FONT
    if record.has_pair and record.is_within_tolerance:
        return _GREEN_FILL, _GREEN_FONT
    if record.has_pair and record.is_overrun:
        return _RED_FILL, _RED_FONT
    return _WHITE_FILL, _BODY_FONT


def _record_delta_fill(record: AuditRecord) -> Optional[PatternFill]:
    """Заливка колонок дельты только для проверенной пары."""
    if not record.has_pair or record.status in {
        STATUS_HITL_PENDING,
        STATUS_REJECTED,
    }:
        return None
    if record.is_within_tolerance:
        return _GREEN_FILL
    if record.is_overrun:
        return _RED_FILL
    return None

def _write_metadata(
    worksheet: Worksheet,
    session: sqlite3.Row,
    generated_at: datetime,
) -> None:
    """Записывает шапку документа: ссылку на дашборд, проект, сессию и дату."""
    worksheet["B1"] = BACK_TO_DASHBOARD_FORMULA
    worksheet["B1"].font = _BACK_LINK_FONT
    worksheet["B1"].fill = _BACK_LINK_FILL
    worksheet["B1"].alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )
    worksheet["C1"] = str(session["project_name"])
    worksheet.merge_cells("C1:Q1")

    worksheet["B2"] = "ID сессии:"
    worksheet["C2"] = str(session["id"])
    worksheet.merge_cells("C2:Q2")

    worksheet["B3"] = "Дата формирования:"
    worksheet["C3"] = generated_at.strftime("%d.%m.%Y %H:%M")
    worksheet.merge_cells("C3:Q3")

    for row in (1, 2, 3):
        if row != 1:
            worksheet[f"B{row}"].font = _META_LABEL_FONT
            worksheet[f"B{row}"].alignment = Alignment(
                horizontal="left", vertical="center"
            )
        worksheet[f"C{row}"].font = _PROJECT_FONT if row == 1 else _META_VALUE_FONT
        worksheet[f"C{row}"].alignment = Alignment(horizontal="left", vertical="center")

    worksheet.row_dimensions[1].height = 24
    worksheet.row_dimensions[2].height = 18
    worksheet.row_dimensions[3].height = 18


def _write_table_header(worksheet: Worksheet) -> None:
    """Формирует строку заголовков таблицы с закреплением под ней."""
    for column, title in COLUMN_TITLES.items():
        cell = worksheet[f"{column}{HEADER_ROW}"]
        cell.value = title
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _CELL_BORDER
    worksheet.row_dimensions[HEADER_ROW].height = 48


def _column_alignment(column: str) -> Alignment:
    """Возвращает выравнивание ячейки в зависимости от типа данных колонки."""
    if column in {"C", "F", "G", "Q"}:
        return Alignment(horizontal="left", vertical="top", wrap_text=True)
    if column in {"A", "B", "D", "H", "P"}:
        return Alignment(horizontal="center", vertical="center", wrap_text=True)
    return Alignment(horizontal="right", vertical="top")


def _apply_number_format(worksheet: Worksheet, row: int) -> None:
    """Назначает числовые форматы колонкам объемов, цен и коэффициентов."""
    for column in ("E", "I", "M", "N"):
        worksheet[f"{column}{row}"].number_format = VOLUME_FORMAT
    for column in ("J", "K", "O"):
        worksheet[f"{column}{row}"].number_format = MONEY_FORMAT
    worksheet[f"L{row}"].number_format = RATIO_FORMAT
    for column in ("A", "B"):
        worksheet[f"{column}{row}"].number_format = INTEGER_FORMAT


def _write_record(
    worksheet: Worksheet,
    row: int,
    index: int,
    record: AuditRecord,
) -> None:
    """Записывает одну строку ведомости и нативные формулы расчета."""
    worksheet[f"A{row}"] = record.match_id
    worksheet[f"B{row}"] = index

    if record.has_rd_pair:
        worksheet[f"C{row}"] = record.rd_description
        worksheet[f"D{row}"] = record.rd_unit
        worksheet[f"E{row}"] = _to_excel_number(record.rd_quantity)

    if record.has_smeta_item:
        worksheet[f"F{row}"] = record.smeta_code
        worksheet[f"G{row}"] = record.smeta_description
        worksheet[f"H{row}"] = record.smeta_unit
        worksheet[f"I{row}"] = _to_excel_number(record.smeta_quantity)
        worksheet[f"J{row}"] = _to_excel_number(record.unit_price)

        if record.smeta_quantity is not None and record.unit_price is not None:
            worksheet[f"K{row}"] = f"=I{row}*J{row}"
        else:
            worksheet[f"K{row}"] = _to_excel_number(record.total_cost)

    if record.has_rd_pair:
        worksheet[f"L{row}"] = _to_excel_number(record.ratio)
        if record.rd_quantity is not None:
            worksheet[f"M{row}"] = f"=E{row}*L{row}"

    if (
        record.has_pair
        and record.status != STATUS_REJECTED
        and record.smeta_quantity is not None
        and record.rd_quantity is not None
    ):
        worksheet[f"N{row}"] = f"=I{row}-M{row}"
        if record.unit_price is not None:
            worksheet[f"O{row}"] = f"=N{row}*J{row}"

    worksheet[f"P{row}"] = record.status
    worksheet[f"Q{row}"] = record.comment

    for column in COLUMN_TITLES:
        cell = worksheet[f"{column}{row}"]
        cell.font = _BODY_FONT
        cell.border = _CELL_BORDER
        cell.alignment = _column_alignment(column)

    _apply_number_format(worksheet, row)

    status_fill, status_font = _record_status_style(record)
    worksheet[f"P{row}"].fill = status_fill
    worksheet[f"P{row}"].font = status_font

    delta_fill = _record_delta_fill(record)
    if delta_fill is not None:
        for column in ("N", "O"):
            worksheet[f"{column}{row}"].fill = delta_fill


def _autosize_columns(worksheet: Worksheet) -> None:
    """Подбирает ширину колонок по контенту с ограничением для описаний."""
    for column, max_width in COLUMN_WIDTH_LIMITS.items():
        longest = len(COLUMN_TITLES[column])
        for row in range(HEADER_ROW, worksheet.max_row + 1):
            value = worksheet[f"{column}{row}"].value
            if value is None:
                continue
            text = str(value)
            if text.startswith("="):
                # Длину формулы Excel отображает результатом, а не текстом.
                continue
            longest = max(longest, *(len(line) for line in text.splitlines()))
        worksheet.column_dimensions[column].width = min(max(longest + 2, 8), max_width)


def _has_price_data(records: Sequence[AuditRecord]) -> bool:
    """Проверяет, есть ли в разделе хотя бы одна ненулевая цена или стоимость."""
    for record in records:
        if not record.has_smeta_item:
            continue
        if record.unit_price not in (None, ZERO):
            return True
        if record.total_cost not in (None, ZERO):
            return True
    return False


def _hide_price_columns(worksheet: Worksheet) -> None:
    """Скрывает визуальный шум из нулевых цен, пока парсер цен не подключен."""
    for column in ("J", "K", "O"):
        worksheet.column_dimensions[column].hidden = True


def _reveal_price_columns(worksheet: Worksheet) -> None:
    """Открывает финансовые колонки, когда в разделе есть реальные цены."""
    for column in ("J", "K", "O"):
        worksheet.column_dimensions[column].hidden = False


def _write_summary_row(
    worksheet: Worksheet,
    records: Sequence[AuditRecord],
    has_prices: bool,
) -> int:
    """Добавляет бухгалтерскую строку «ИТОГО ПО РАЗДЕЛУ» с формулами SUM.

    Объемные итоги (``E``, ``I``, ``M``, ``N``) выводятся всегда. Финансовые
    итоги (``K`` — стоимость по смете, ``O`` — дельта стоимости) выводятся
    только когда в разделе есть реальные цены, иначе колонки скрыты.

    :return: номер итоговой строки.
    """
    summary_row = FIRST_DATA_ROW + len(records)
    last_data_row = summary_row - 1

    worksheet[f"B{summary_row}"] = "ИТОГО ПО РАЗДЕЛУ:"
    for column in ("E", "I", "M", "N"):
        cell = worksheet[f"{column}{summary_row}"]
        if records:
            cell.value = f"=SUM({column}{FIRST_DATA_ROW}:{column}{last_data_row})"
        else:
            cell.value = 0
        cell.number_format = VOLUME_FORMAT

    if has_prices:
        for column in ("K", "O"):
            cell = worksheet[f"{column}{summary_row}"]
            if records:
                cell.value = (
                    f"=SUM({column}{FIRST_DATA_ROW}:{column}{last_data_row})"
                )
            else:
                cell.value = 0
            cell.number_format = MONEY_FORMAT

    for column in range(2, len(COLUMN_TITLES) + 1):
        cell = worksheet.cell(row=summary_row, column=column)
        cell.font = _SUMMARY_FONT
        cell.fill = _SUMMARY_FILL
        cell.border = _SUMMARY_BORDER
        if column == 2:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        elif column in {5, 9, 11, 13, 14, 15}:
            cell.alignment = Alignment(horizontal="right", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.row_dimensions[summary_row].height = 22
    return summary_row


def _populate_worksheet(
    worksheet: Worksheet,
    session: sqlite3.Row,
    records: Sequence[AuditRecord],
    generated_at: datetime,
) -> None:
    """Заполняет один лист ведомости шапкой, таблицей и формулами."""
    worksheet.column_dimensions[HIDDEN_ID_COLUMN].hidden = True
    _write_metadata(worksheet, session, generated_at)
    _write_table_header(worksheet)

    for offset, record in enumerate(records):
        _write_record(worksheet, FIRST_DATA_ROW + offset, offset + 1, record)

    # Финансовые колонки активируются только при наличии реальных цен:
    # иначе остаются скрытыми как визуальный шум нулевых расценок.
    has_prices = _has_price_data(records)
    if has_prices:
        _reveal_price_columns(worksheet)
    else:
        _hide_price_columns(worksheet)

    last_data_row = FIRST_DATA_ROW + len(records) - 1
    _write_summary_row(worksheet, records, has_prices)

    _autosize_columns(worksheet)
    # Плашка возврата должна быть читаемой целиком, поэтому колонка B
    # расширяется после авторазмера (для данных она остается узкой).
    worksheet.column_dimensions["B"].width = max(
        worksheet.column_dimensions["B"].width or 0, 20
    )
    worksheet.auto_filter.ref = (
        f"B{HEADER_ROW}:Q{max(last_data_row, HEADER_ROW)}"
    )
    worksheet.freeze_panes = f"C{FIRST_DATA_ROW}"
    worksheet.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    worksheet.sheet_view.showGridLines = False


_INVALID_SHEET_TITLE = re.compile(r"[\\/*?:\[\]]")
_SECTION_CODE_PATTERN = re.compile(r"(?<!\d)(\d{2}-\d{2}-\d{2})(?!\d)")
_LEADING_LS_PATTERN = re.compile(r"^\s*лс\s*", re.IGNORECASE)


def _sanitize_sheet_title(raw: Optional[str], fallback: str) -> str:
    """Готовит допустимое имя листа Excel (не более 31 символа)."""
    title = _INVALID_SHEET_TITLE.sub(" ", raw or "")
    title = re.sub(r"\s+", " ", title).strip().strip("'")[:31].strip()
    if title:
        return title
    safe_fallback = _INVALID_SHEET_TITLE.sub(" ", fallback)
    return re.sub(r"\s+", " ", safe_fallback).strip()[:31] or "Раздел"


def _extract_section_code(raw: Optional[str]) -> Optional[str]:
    """Извлекает шифр раздела формата ``NN-NN-NN`` из имени листа сметы."""
    match = _SECTION_CODE_PATTERN.search(raw or "")
    return match.group(1) if match else None


def _section_label_and_cipher(
    raw: Optional[str],
    fallback: str,
) -> tuple[str, str]:
    """Возвращает читаемое имя раздела и его шифр для вкладки/дашборда."""
    code = _extract_section_code(raw)
    if code is None:
        label = _sanitize_sheet_title(raw, fallback)
        return label, "—"

    cipher = f"ЛС {code}"
    mapped = SECTION_NAMES_BY_CODE.get(code)
    if mapped is not None:
        return mapped, cipher

    source = _LEADING_LS_PATTERN.sub("", (raw or "").strip())
    prefix = source.split(code, maxsplit=1)[0].strip(" -–—()[]")
    if prefix:
        return _sanitize_sheet_title(prefix, fallback), cipher
    return f"Раздел {code}", cipher


def _readable_sheet_title(raw: Optional[str], fallback: str) -> str:
    """Формирует имя вкладки ``Раздел (ЛС NN-NN-NN)`` в пределах лимита Excel."""
    label, cipher = _section_label_and_cipher(raw, fallback)
    if cipher == "—":
        return _sanitize_sheet_title(label, fallback)
    if label.startswith("Раздел "):
        return _sanitize_sheet_title(cipher, fallback)
    return _sanitize_sheet_title(f"{label} ({cipher})", fallback)


def _unique_sheet_title(base: str, used: set[str]) -> str:
    """Возвращает уникальное имя листа с учетом регистронезависимых коллизий."""
    title = base
    suffix = 2
    while title.casefold() in used:
        suffix_text = f" ({suffix})"
        title = f"{base[: 31 - len(suffix_text)].rstrip()}{suffix_text}"
        suffix += 1
    used.add(title.casefold())
    return title


# ----------------------------------------------------------------------
# 4. Сводный дашборд и метрики разделов
# ----------------------------------------------------------------------


@dataclass(frozen=True)
class SectionExport:
    """Подготовленная к выгрузке сессия одного раздела."""

    session: sqlite3.Row
    records: tuple[AuditRecord, ...]
    title: str
    label: str
    cipher: str


@dataclass(frozen=True)
class SectionMetrics:
    """Управленческие метрики одного раздела для сводного дашборда."""

    session_id: str
    title: str
    label: str
    cipher: str
    total: int
    smeta_total: int
    rd_total: int
    equipment_total: int
    equipment_paired: int
    smr_total: int
    unclassified_total: int
    paired: int
    no_rd: int
    no_smeta: int
    discrepancies: int
    pending: int
    rejected: int
    incomplete: int
    total_cost: Decimal
    delta_cost: Decimal
    has_prices: bool

    @property
    def coverage(self) -> float:
        """Покрытие спецификациями РД: пары / оборудование и материалы."""
        if not self.equipment_total:
            return 0.0
        return self.equipment_paired / self.equipment_total

    @property
    def verdict(self) -> str:
        """Краткий управленческий вердикт по разделу."""
        if self.total == 0:
            return VERDICT_NO_DATA
        if self.paired == 0:
            if self.rd_total == 0:
                return VERDICT_NO_RD_SPEC
            return VERDICT_NO_MATCH
        if self.discrepancies > 0:
            return VERDICT_DISCREPANCY
        if (
            self.no_rd
            or self.no_smeta
            or self.pending
            or self.rejected
            or self.incomplete
            or self.total != self.paired + self.no_rd + self.no_smeta
        ):
            return VERDICT_REVIEW
        return VERDICT_OK

    @property
    def fully_verified(self) -> bool:
        """Раздел проверен полностью и не требует внимания эксперта."""
        return self.verdict == VERDICT_OK


def _measure_section(section: SectionExport) -> SectionMetrics:
    """Считает KPI раздела по нормализованным строкам ведомости."""
    paid_records = [record for record in section.records if record.has_pair]
    total_cost = sum(
        (
            record.estimated_cost
            for record in section.records
            if record.has_smeta_item and record.estimated_cost is not None
        ),
        ZERO,
    )
    delta_cost = sum(
        (
            record.delta_volume * record.unit_price
            for record in paid_records
            if record.status != STATUS_REJECTED
            and record.delta_volume is not None
            and record.unit_price is not None
        ),
        ZERO,
    )
    discrepancies = sum(
        1
        for record in paid_records
        if record.status != STATUS_REJECTED
        and record.delta_volume is not None
        and not record.is_within_tolerance
    )
    incomplete = sum(
        1
        for record in paid_records
        if record.status != STATUS_REJECTED and record.delta_volume is None
    )
    no_rd = sum(
        1
        for record in section.records
        if record.has_smeta_item and not record.has_rd_pair
    )
    no_smeta = sum(
        1
        for record in section.records
        if record.has_rd_pair and not record.has_smeta_item
    )
    equipment_records = [
        record
        for record in section.records
        if record.has_smeta_item and record.smeta_kind == KIND_EQUIPMENT
    ]
    smeta_records = [
        record for record in section.records if record.has_smeta_item
    ]

    return SectionMetrics(
        session_id=str(section.session["id"]),
        title=section.title,
        label=section.label,
        cipher=section.cipher,
        total=len(section.records),
        smeta_total=len(smeta_records),
        rd_total=sum(1 for record in section.records if record.has_rd_pair),
        equipment_total=len(equipment_records),
        equipment_paired=sum(
            1 for record in equipment_records if record.has_pair
        ),
        smr_total=sum(
            1
            for record in smeta_records
            if record.smeta_kind == KIND_SMR
        ),
        unclassified_total=sum(
            1
            for record in smeta_records
            if record.smeta_kind == KIND_UNCLASSIFIED
        ),
        paired=len(paid_records),
        no_rd=no_rd,
        no_smeta=no_smeta,
        discrepancies=discrepancies,
        pending=sum(
            1 for record in section.records if record.status == STATUS_HITL_PENDING
        ),
        rejected=sum(
            1 for record in section.records if record.status == STATUS_REJECTED
        ),
        incomplete=incomplete,
        total_cost=total_cost,
        delta_cost=delta_cost,
        has_prices=_has_price_data(section.records),
    )


def _verdict_style(verdict: str) -> tuple[PatternFill, Font]:
    """Возвращает стиль вердикта для дашборда и цвет вкладки раздела."""
    if verdict == VERDICT_OK:
        return _GREEN_FILL, _GREEN_FONT
    if verdict == VERDICT_DISCREPANCY:
        return _RED_FILL, _RED_FONT
    if verdict == VERDICT_REVIEW:
        return _YELLOW_FILL, _YELLOW_FONT
    if verdict in {VERDICT_NO_MATCH, VERDICT_NO_RD_SPEC}:
        return _WARNING_FILL, _WARNING_FONT
    if verdict == VERDICT_NO_DATA:
        return _GRAY_FILL, _GRAY_FONT
    return _WHITE_FILL, _BODY_FONT


_TAB_COLORS: Final[dict[str, str]] = {
    VERDICT_OK: "00B050",
    VERDICT_DISCREPANCY: "C00000",
    VERDICT_REVIEW: "FFC000",
    VERDICT_NO_MATCH: "FFC000",
    VERDICT_NO_RD_SPEC: "FFC000",
    VERDICT_NO_DATA: "808080",
}


_DASHBOARD_LAST_COLUMN: Final[int] = 13
_DASHBOARD_TABLE_TITLES: Final[dict[str, str]] = {
    "A": "Раздел",
    "B": "Шифр",
    "C": "Позиций сметы",
    "D": "Оборудование",
    "E": "Сопоставлено РД",
    "F": "Покрытие РД",
    "G": "СМР",
    "H": "Расхождения",
    "I": "На проверке",
    "J": "Стоимость сметы, руб.",
    "K": "Дельта стоимости, руб.",
    "L": "Вердикт",
}
_DASHBOARD_WIDTHS: Final[dict[str, int]] = {
    "A": 28,
    "B": 15,
    "C": 10,
    "D": 12,
    "E": 14,
    "F": 11,
    "G": 9,
    "H": 13,
    "I": 12,
    "J": 21,
    "K": 23,
    "L": 38,
    "M": 3,
}

_DASHBOARD_TITLE_FONT: Final[Font] = Font(
    name="Calibri", size=20, bold=True, color="FFFFFF"
)
_DASHBOARD_TITLE_FILL: Final[PatternFill] = PatternFill(
    "solid", fgColor="1F4E78"
)
_DASHBOARD_SUBTITLE_FONT: Final[Font] = Font(
    name="Calibri", size=14, bold=True, color="1F4E78"
)
_DASHBOARD_META_FONT: Final[Font] = Font(
    name="Calibri", size=10, color="595959"
)
_DASHBOARD_NOTE_FILL: Final[PatternFill] = PatternFill(
    "solid", fgColor="FFF2CC"
)
_DASHBOARD_NOTE_FONT: Final[Font] = Font(
    name="Calibri", size=10, italic=True, color="7F6000"
)
_DASHBOARD_KPI_LABEL_FILL: Final[PatternFill] = PatternFill(
    "solid", fgColor="DDEBF7"
)
_DASHBOARD_KPI_VALUE_FONT: Final[Font] = Font(
    name="Calibri", size=12, bold=True, color="1F1F1F"
)
_DASHBOARD_SECTION_FONT: Final[Font] = Font(
    name="Calibri", size=12, bold=True, color="FFFFFF"
)
_DASHBOARD_LINK_FONT: Final[Font] = Font(
    name="Calibri", size=10, color="0563C1", underline="single"
)
_DASHBOARD_ALT_FILL: Final[PatternFill] = PatternFill(
    "solid", fgColor="F7F9FC"
)


def _style_dashboard_row(
    worksheet: Worksheet,
    row: int,
    fill: PatternFill,
) -> None:
    """Заливает всю строку дашборда, включая объединенные ячейки."""
    for column in range(1, _DASHBOARD_LAST_COLUMN + 1):
        worksheet.cell(row=row, column=column).fill = fill


def _write_dashboard_banner(
    worksheet: Worksheet,
    project_label: str,
    generated_at: datetime,
) -> None:
    """Формирует шапку титульного листа."""
    worksheet.merge_cells("A1:M1")
    worksheet["A1"] = "СВОДНЫЙ АУДИТ ОБЪЁМОВ"
    worksheet["A1"].font = _DASHBOARD_TITLE_FONT
    worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    _style_dashboard_row(worksheet, 1, _DASHBOARD_TITLE_FILL)
    worksheet.row_dimensions[1].height = 36

    worksheet.merge_cells("A2:M2")
    worksheet["A2"] = project_label
    worksheet["A2"].font = _DASHBOARD_SUBTITLE_FONT
    worksheet["A2"].alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[2].height = 24

    worksheet.merge_cells("A3:M3")
    worksheet["A3"] = (
        f"Сформировано: {generated_at.strftime('%d.%m.%Y %H:%M')}  |  "
        "Источник данных: SQLite (SSOT)"
    )
    worksheet["A3"].font = _DASHBOARD_META_FONT
    worksheet["A3"].alignment = Alignment(horizontal="left", vertical="center")

    worksheet.merge_cells("A4:M4")
    worksheet["A4"] = (
        "Методика: покрытие считается по оборудованию и материалам (главы 1, 2, "
        "13 и «Прайс-лист»). Монтажные и строительные работы (главы 3, 4, 6) — "
        "норматив СМР и не имеют отдельных строк в спецификациях РД."
    )
    worksheet["A4"].font = _DASHBOARD_NOTE_FONT
    worksheet["A4"].alignment = Alignment(
        horizontal="left", vertical="center", wrap_text=True
    )
    _style_dashboard_row(worksheet, 4, _DASHBOARD_NOTE_FILL)
    worksheet.row_dimensions[4].height = 32


def _write_dashboard_kpi(
    worksheet: Worksheet,
    row: int,
    label: str,
    value: Any,
    note: str,
    *,
    number_format: Optional[str] = None,
    value_fill: Optional[PatternFill] = None,
) -> None:
    """Записывает одну KPI-строку дашборда."""
    label_cell = worksheet.cell(row=row, column=1, value=label)
    value_cell = worksheet.cell(row=row, column=2, value=value)
    note_cell = worksheet.cell(row=row, column=3, value=note)
    worksheet.merge_cells(
        start_row=row, start_column=3, end_row=row, end_column=12
    )

    label_cell.font = _META_LABEL_FONT
    label_cell.fill = _DASHBOARD_KPI_LABEL_FILL
    label_cell.alignment = Alignment(
        horizontal="left", vertical="center", wrap_text=True
    )
    value_cell.font = _DASHBOARD_KPI_VALUE_FONT
    value_cell.alignment = Alignment(
        horizontal="center" if isinstance(value, str) else "right",
        vertical="center",
    )
    note_cell.font = _DASHBOARD_META_FONT
    note_cell.alignment = Alignment(
        horizontal="left", vertical="center", wrap_text=True
    )

    if number_format is not None:
        value_cell.number_format = number_format
    if value_fill is not None:
        value_cell.fill = value_fill

    for column in range(1, 13):
        worksheet.cell(row=row, column=column).border = _CELL_BORDER
    worksheet.row_dimensions[row].height = 26 if len(label) > 28 else 21


def _write_dashboard_kpis(
    worksheet: Worksheet,
    metrics: Sequence[SectionMetrics],
) -> None:
    """Выводит ключевые показатели объекта."""
    total_positions = sum(metric.smeta_total for metric in metrics)
    equipment_total = sum(metric.equipment_total for metric in metrics)
    equipment_paired = sum(metric.equipment_paired for metric in metrics)
    smr_total = sum(metric.smr_total for metric in metrics)
    unclassified_total = sum(
        metric.unclassified_total for metric in metrics
    )
    discrepancies = sum(metric.discrepancies for metric in metrics)
    pending = sum(metric.pending for metric in metrics)
    coverage = (
        equipment_paired / equipment_total if equipment_total else 0.0
    )
    has_prices = any(metric.has_prices for metric in metrics)
    total_cost: Any = (
        _to_excel_number(sum((m.total_cost for m in metrics), ZERO))
        if has_prices
        else "—"
    )
    delta_cost: Any = (
        _to_excel_number(sum((m.delta_cost for m in metrics), ZERO))
        if has_prices
        else "—"
    )
    coverage_fill = (
        _GREEN_FILL
        if equipment_total and abs(coverage - 1.0) < 1e-9
        else None
    )

    worksheet.merge_cells("A5:M5")
    worksheet["A5"] = "КЛЮЧЕВЫЕ ПОКАЗАТЕЛИ ОБЪЕКТА"
    worksheet["A5"].font = _META_LABEL_FONT
    worksheet["A5"].alignment = Alignment(horizontal="left", vertical="center")
    _style_dashboard_row(worksheet, 5, _DASHBOARD_KPI_LABEL_FILL)

    for column, title in (
        ("A", "Показатель"),
        ("B", "Значение"),
        ("C", "Комментарий"),
    ):
        cell = worksheet[f"{column}{DASHBOARD_KPI_HEADER_ROW}"]
        cell.value = title
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _CELL_BORDER
    worksheet.merge_cells(
        f"C{DASHBOARD_KPI_HEADER_ROW}:L{DASHBOARD_KPI_HEADER_ROW}"
    )

    rows: list[
        tuple[str, Any, str, Optional[str], Optional[PatternFill]]
    ] = [
        (
            "Всего позиций сметы",
            total_positions,
            "Строк сметы в ведомостях разделов",
            INTEGER_FORMAT,
            None,
        ),
        (
            "Оборудование и материалы (в РД)",
            equipment_total,
            "Главы 1, 2, 13 и строки «Прайс-лист»",
            INTEGER_FORMAT,
            None,
        ),
        (
            "Покрытие спецификациями РД",
            coverage,
            "Сопоставлено / Оборудование и материалы",
            PERCENT_FORMAT,
            coverage_fill,
        ),
        (
            "Строительно-монтажные работы (норматив СМР)",
            smr_total,
            "Технологические работы монтажа без отдельных строк "
            "в спецификациях",
            INTEGER_FORMAT,
            None,
        ),
        (
            "Выявлено расхождений объёмов",
            discrepancies,
            "Отклонения сверх допуска, включая превышения",
            INTEGER_FORMAT,
            _RED_FILL if discrepancies else None,
        ),
        (
            "Очередь ручной проверки (HITL)",
            pending,
            "Позиции, требующие решения эксперта",
            INTEGER_FORMAT,
            _YELLOW_FILL if pending else None,
        ),
    ]
    if unclassified_total:
        rows.append(
            (
                "Позиции вне классификации",
                unclassified_total,
                "Шифры вне глав 1, 2, 13 и 3, 4, 6",
                INTEGER_FORMAT,
                _YELLOW_FILL,
            )
        )
    rows.extend(
        [
            (
                "Стоимость сметы, руб.",
                total_cost,
                "Цены не подключены"
                if not has_prices
                else "Сумма по сметным позициям",
                MONEY_FORMAT if has_prices else None,
                None,
            ),
            (
                "Дельта стоимости, руб.",
                delta_cost,
                "Смета минус приведенный объем РД",
                MONEY_FORMAT if has_prices else None,
                None,
            ),
        ]
    )

    for offset, (label, value, note, number_format, fill) in enumerate(rows):
        _write_dashboard_kpi(
            worksheet,
            DASHBOARD_FIRST_KPI_ROW + offset,
            label,
            value,
            note,
            number_format=number_format,
            value_fill=fill,
        )


def _write_dashboard_section_table(
    worksheet: Worksheet,
    metrics: Sequence[SectionMetrics],
) -> None:
    """Выводит таблицу разделов с автофильтром и переходами на вкладки."""
    worksheet.merge_cells(
        f"A{DASHBOARD_SECTION_TITLE_ROW}:L{DASHBOARD_SECTION_TITLE_ROW}"
    )
    title_cell = worksheet[f"A{DASHBOARD_SECTION_TITLE_ROW}"]
    title_cell.value = "РАЗДЕЛЫ ОБЪЕКТА"
    title_cell.font = _DASHBOARD_SECTION_FONT
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    _style_dashboard_row(
        worksheet,
        DASHBOARD_SECTION_TITLE_ROW,
        _DASHBOARD_TITLE_FILL,
    )
    worksheet.row_dimensions[DASHBOARD_SECTION_TITLE_ROW].height = 24

    for column, title in _DASHBOARD_TABLE_TITLES.items():
        cell = worksheet[f"{column}{DASHBOARD_SECTION_HEADER_ROW}"]
        cell.value = title
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _CELL_BORDER
    worksheet.row_dimensions[DASHBOARD_SECTION_HEADER_ROW].height = 40

    for offset, metric in enumerate(metrics):
        row = DASHBOARD_FIRST_SECTION_ROW + offset
        values: dict[str, Any] = {
            "A": metric.label,
            "B": metric.cipher,
            "C": metric.smeta_total,
            "D": metric.equipment_total,
            "E": metric.equipment_paired,
            "F": metric.coverage,
            "G": metric.smr_total,
            "H": metric.discrepancies,
            "I": metric.pending,
            "J": (
                _to_excel_number(metric.total_cost)
                if metric.has_prices
                else "—"
            ),
            "K": (
                _to_excel_number(metric.delta_cost)
                if metric.has_prices
                else "—"
            ),
            "L": metric.verdict,
        }
        row_fill = _DASHBOARD_ALT_FILL if offset % 2 else _WHITE_FILL
        for column, value in values.items():
            cell = worksheet[f"{column}{row}"]
            cell.value = value
            cell.font = _BODY_FONT
            cell.fill = row_fill
            cell.border = _CELL_BORDER
            if column == "A":
                cell.alignment = Alignment(
                    horizontal="left", vertical="center", wrap_text=True
                )
            elif column == "B":
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            elif column == "L":
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            elif column in {"J", "K"} and isinstance(value, str):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")

        for column in ("C", "D", "E", "G", "H", "I"):
            worksheet[f"{column}{row}"].number_format = INTEGER_FORMAT
        worksheet[f"F{row}"].number_format = PERCENT_FORMAT
        if metric.has_prices:
            for column in ("J", "K"):
                worksheet[f"{column}{row}"].number_format = MONEY_FORMAT

        link_cell = worksheet[f"A{row}"]
        escaped_title = metric.title.replace("'", "''")
        link_cell.hyperlink = f"#'{escaped_title}'!A1"
        link_cell.font = _DASHBOARD_LINK_FONT

        if metric.discrepancies:
            worksheet[f"H{row}"].fill = _RED_FILL
            worksheet[f"H{row}"].font = _RED_FONT
        if metric.pending:
            worksheet[f"I{row}"].fill = _YELLOW_FILL
            worksheet[f"I{row}"].font = _YELLOW_FONT
        verdict_fill, verdict_font = _verdict_style(metric.verdict)
        worksheet[f"L{row}"].fill = verdict_fill
        worksheet[f"L{row}"].font = verdict_font
        worksheet.row_dimensions[row].height = (
            42
            if metric.verdict in {VERDICT_NO_MATCH, VERDICT_NO_RD_SPEC}
            else 24
        )

    last_row = max(
        DASHBOARD_SECTION_HEADER_ROW,
        DASHBOARD_FIRST_SECTION_ROW + len(metrics) - 1,
    )
    worksheet.auto_filter.ref = (
        f"A{DASHBOARD_SECTION_HEADER_ROW}:L{last_row}"
    )


def _project_label(sections: Sequence[SectionExport]) -> str:
    """Формирует подпись проекта для дашборда."""
    names = sorted({str(section.session["project_name"]) for section in sections})
    if not names:
        return "Проект не указан"
    if len(names) == 1:
        return names[0]
    return f"Несколько проектов ({len(names)}): " + "; ".join(names)


def _render_dashboard(
    worksheet: Worksheet,
    sections: Sequence[SectionExport],
    metrics: Sequence[SectionMetrics],
    generated_at: datetime,
) -> None:
    """Полностью оформляет титульный лист книги."""
    _write_dashboard_banner(worksheet, _project_label(sections), generated_at)
    _write_dashboard_kpis(worksheet, metrics)
    _write_dashboard_section_table(worksheet, metrics)

    for column, width in _DASHBOARD_WIDTHS.items():
        worksheet.column_dimensions[column].width = width
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 90
    worksheet.sheet_properties.tabColor = "1F4E78"
    worksheet.page_setup.orientation = "landscape"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_area = (
        f"A1:L{DASHBOARD_FIRST_SECTION_ROW + len(metrics) - 1}"
    )


# ----------------------------------------------------------------------
# 5. Публичный API и точка входа CLI
# ----------------------------------------------------------------------


def _session_fallback(session: sqlite3.Row) -> str:
    """Возвращает запасное имя раздела из проекта и имени тома РД."""
    project_name = str(session["project_name"])
    if not session["smeta_sheet"]:
        return project_name
    return f"{project_name} {Path(session['rd_filename']).stem}"


def _section_export(
    session: sqlite3.Row,
    records: Sequence[AuditRecord],
    title: str,
) -> SectionExport:
    """Создает описание раздела для дашборда и выгрузки."""
    fallback = _session_fallback(session)
    label, cipher = _section_label_and_cipher(session["smeta_sheet"], fallback)
    return SectionExport(
        session=session,
        records=tuple(records),
        title=title,
        label=label,
        cipher=cipher,
    )


def _prepare_sections(
    conn: sqlite3.Connection,
    session_ids: Sequence[str],
) -> list[SectionExport]:
    """Читает сессии и назначает читаемые уникальные имена вкладок."""
    used_titles: set[str] = {DASHBOARD_TITLE.casefold()}
    sections: list[SectionExport] = []
    for session_id in session_ids:
        resolved_id = resolve_session_id(conn, session_id)
        session = _load_session(conn, resolved_id)
        records = load_audit_records(conn, resolved_id)
        fallback = _session_fallback(session)
        base_title = _readable_sheet_title(session["smeta_sheet"], fallback)
        title = _unique_sheet_title(base_title, used_titles)
        sections.append(_section_export(session, records, title))
    return sections


def _build_workbook(
    sections: Sequence[SectionExport],
    generated_at: datetime,
) -> Workbook:
    """Собирает книгу: дашборд первым листом, затем разделы."""
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = DASHBOARD_TITLE
    metrics = [_measure_section(section) for section in sections]

    for section, metric in zip(sections, metrics):
        worksheet = workbook.create_sheet(title=section.title)
        _populate_worksheet(
            worksheet,
            section.session,
            section.records,
            generated_at,
        )
        tab_color = _TAB_COLORS.get(metric.verdict)
        if tab_color is not None:
            worksheet.sheet_properties.tabColor = tab_color

    _render_dashboard(dashboard, sections, metrics, generated_at)
    workbook.active = 0
    # Excel пересчитывает нативные формулы при открытии книги.
    workbook.calculation.fullCalcOnLoad = True
    return workbook


def _save_workbook(workbook: Workbook, output_path: str | Path) -> Path:
    """Сохраняет книгу на диск, создавая родительские каталоги."""
    workbook.properties.creator = "sst-audit"
    workbook.properties.title = "Сводный аудит объёмов"
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return target


def export_audit_to_excel(
    conn: sqlite3.Connection,
    session_id: str,
    output_path: str | Path,
) -> Path:
    """Формирует книгу с дашбордом и ведомостью одного раздела.

    :param conn: соединение с SQLite (SSOT).
    :param session_id: идентификатор сессии либо ``latest``.
    :param output_path: путь к создаваемому файлу отчета.
    :return: путь к сохраненному файлу.
    :raises LookupError: сессия не найдена.
    """
    resolved_id = resolve_session_id(conn, session_id)
    session = _load_session(conn, resolved_id)
    records = load_audit_records(conn, resolved_id)
    # Одиночный запуск получает осмысленное имя раздела, например
    # «СОТ (ЛС 02-01-14)», а при отсутствии шифра — имя проекта.
    sheet_title = _readable_sheet_title(
        session["smeta_sheet"], str(session["project_name"])
    )
    section = _section_export(session, records, sheet_title)

    workbook = _build_workbook([section], datetime.now())
    return _save_workbook(workbook, output_path)


def export_sessions_to_excel(
    conn: sqlite3.Connection,
    session_ids: Sequence[str],
    output_path: str | Path,
) -> Path:
    """Формирует многолистовую книгу: дашборд и листы разделов.

    Используется ночным пакетным режимом, когда каждая сессия соответствует
    отдельному тому РД и листу ``ЛС ...``. Имена вкладок формируются как
    ``АР (ЛС 02-01-01)``; коллизии и недопустимые символы Excel разрешаются
    автоматически.

    :param conn: соединение с SQLite (SSOT).
    :param session_ids: идентификаторы успешно завершенных сессий.
    :param output_path: путь к создаваемому XLSX-отчету.
    :return: путь к сохраненному файлу.
    :raises ValueError: не передано ни одной сессии.
    :raises LookupError: одна из сессий не найдена.
    """
    ordered_ids = list(session_ids)
    if not ordered_ids:
        raise ValueError("Для многолистового экспорта не передано ни одной сессии")

    sections = _prepare_sections(conn, ordered_ids)
    workbook = _build_workbook(sections, datetime.now())
    return _save_workbook(workbook, output_path)


def build_parser() -> argparse.ArgumentParser:
    """Создает парсер аргументов командной строки экспортера."""
    parser = argparse.ArgumentParser(
        prog="exporter.py",
        description=(
            "Формирование ведомости аудита объемов в Excel "
            "с нативными формулами проверки"
        ),
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument(
        "--db",
        dest="db_path",
        default=db.DEFAULT_DB_PATH,
        help="Путь к файлу SQLite (SSOT)",
    )
    parser.add_argument(
        "--session",
        default=LATEST_SESSION,
        help="Идентификатор сессии либо 'latest' для последней созданной",
    )
    parser.add_argument(
        "--out",
        dest="output_path",
        default=DEFAULT_OUTPUT,
        help="Путь к создаваемому XLSX-отчету",
    )
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    """Запускает экспорт ведомости; возвращает код выхода процесса."""
    args = build_parser().parse_args(argv)
    db_path = Path(args.db_path)

    if not db_path.is_file():
        print(f"Ошибка: файл БД не найден: {db_path}")
        return 2

    conn: Optional[sqlite3.Connection] = None
    try:
        conn = db.get_connection(db_path)
        db.init_db(conn)
        try:
            resolved_id = resolve_session_id(conn, args.session)
            target = export_audit_to_excel(conn, resolved_id, args.output_path)
            row_count = len(load_audit_records(conn, resolved_id))
        except LookupError as exc:
            print(f"Ошибка: {exc}")
            return 2
    except sqlite3.Error as exc:
        print(f"Ошибка SQLite: {exc}")
        return 1
    except OSError as exc:
        print(f"Не удалось сохранить отчет: {exc}")
        return 1
    finally:
        if conn is not None:
            conn.close()

    print(f"Ведомость сформирована: {target} (строк: {row_count})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())




